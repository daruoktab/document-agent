"""Regression tests for native and visual Excel region extraction."""

import gc
import sqlite3
import tempfile
import unittest
from contextlib import closing
from pathlib import Path

from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.table import Table

from app.config import Settings
from app.excel import (
    ExcelTileBudgetError,
    _convert_excel_regions_to_pdf,
    _convert_workbook_to_xlsx_copy,
    _estimate_native_tokens,
    compose_excel_markdown,
    persist_excel_native_data,
    plan_excel_visual_regions,
    split_excel_regions_into_tiles,
    survey_excel_workbook,
)
from app.pdf import pdf_page_count
from app.ppt import _find_libreoffice_binary
from app.schemas import ExcelNativeArtifact


class TestExcelRegionSurvey(unittest.TestCase):
    def _build_adaptive_workbook(self, path: Path) -> None:
        workbook = Workbook()
        data = workbook.active
        data.title = "Records"
        data.append(["ID", "Area", "Amount", "Status"])
        for row in range(1, 41):
            data.append([row, f"Area {row % 5}", row * 1250, "Open"])

        summary = workbook.create_sheet("Rollup")
        summary.append(["Metric", "Value"])
        for row in range(2, 14):
            summary.append([f"Metric {row - 1}", f"=SUM(Records!C2:C{row + 1})"])

        support = workbook.create_sheet("ChartData")
        support.append(["Period", "Value"])
        for row in range(2, 14):
            support.append([f"Period {row - 1}", f"=Records!C{row}"])

        dashboard = workbook.create_sheet("Visual")
        dashboard.merge_cells("A1:H1")
        dashboard["A1"] = "Operational Monitoring"
        dashboard["A3"] = "=ChartData!A1"
        chart = LineChart()
        chart.title = "Movement"
        chart.add_data(
            Reference(support, min_col=2, min_row=1, max_row=13),
            titles_from_data=True,
        )
        chart.set_categories(Reference(support, min_col=1, min_row=2, max_row=13))
        chart.anchor = "A5"
        dashboard.add_chart(chart)
        workbook.save(path)

    def test_token_budget_splits_dense_cells_without_losing_evidence(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "dense.xlsx"
            workbook = Workbook()
            sheet = workbook.active
            for row in range(1, 15):
                sheet.cell(row, 1, f"Row {row}")
                sheet.cell(row, 2, "=SUM(" + ",".join(["12345"] * 30) + ")")
            workbook.save(path)
            survey = survey_excel_workbook(path)
            settings = Settings(excel_tile_max_native_tokens=450)
            tiled = split_excel_regions_into_tiles(survey, settings)
            regions = tiled.sheets[0].regions
            self.assertGreater(len(regions), 1)
            self.assertTrue(all(_estimate_native_tokens(r.native_text) <= 450 for r in regions))
            expected = {c.coordinate: c.formula for r in survey.sheets[0].regions for c in r.cells}
            actual = {c.coordinate: c.formula for r in regions for c in r.cells}
            self.assertEqual(actual, expected)
            with self.assertRaises(ExcelTileBudgetError):
                split_excel_regions_into_tiles(survey, Settings(excel_tile_max_native_tokens=10))

    def test_merged_caption_does_not_bridge_tables(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "caption.xlsx"
            workbook = Workbook()
            sheet = workbook.active
            sheet.merge_cells("A1:N1")
            sheet["A1"] = "Shared caption"
            for column, marker in ((1, "LEFT"), (9, "RIGHT")):
                for row in range(2, 40):
                    for offset in range(4):
                        sheet.cell(row, column + offset, f"{marker}{row}_{offset}")
            workbook.save(path)
            survey = survey_excel_workbook(path)
            regions = survey.sheets[0].regions
            self.assertEqual([r.cell_range for r in regions], ["A1:N1", "A2:D39", "I2:L39"])
            self.assertNotIn("RIGHT", regions[1].native_text)
            self.assertNotIn("LEFT", regions[2].native_text)

    @unittest.skipUnless(_find_libreoffice_binary(), "LibreOffice tidak tersedia")
    def test_vertical_tiles_render_only_their_rows(self) -> None:
        import pymupdf

        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            path = root / "tall.xlsx"
            workbook = Workbook()
            sheet = workbook.active
            sheet.column_dimensions["A"].width = 24
            for row in range(1, 45):
                sheet.cell(row, 1, f"MARKER{row:03d}")
                sheet.cell(row, 2, row)
            workbook.save(path)
            settings = Settings(excel_tile_max_columns=8, excel_tile_max_rows=20)
            survey = split_excel_regions_into_tiles(
                survey_excel_workbook(path, settings=settings), settings,
            )
            regions = survey.sheets[0].regions
            pdf = root / "tiles.pdf"
            self.assertTrue(_convert_excel_regions_to_pdf(
                path, pdf, root / "profile", regions,
            ))
            with pymupdf.open(pdf) as document:
                self.assertEqual(len(document), len(regions))
                for page, region in zip(document, regions, strict=True):
                    text = page.get_text()
                    expected_rows = {cell.row for cell in region.cells}
                    for row in range(1, 45):
                        self.assertEqual(
                            f"MARKER{row:03d}" in text, row in expected_rows,
                            f"{region.region_id}: wrong presence of row {row}",
                        )

    def _build_side_by_side_workbook(self, path: Path) -> None:
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Billing"

        for start_column, month, values in (
            (2, "FEB 2026", (3287, 1_036_330_654, 113_993_234, 1_150_323_888)),
            (8, "MAR 2026", (3035, 951_306_917, 104_641_630, 1_055_948_547)),
        ):
            sheet.merge_cells(
                start_row=2,
                start_column=start_column,
                end_row=2,
                end_column=start_column + 4,
            )
            title = sheet.cell(2, start_column, f"INVOICE 1 {month}")
            title.font = Font(bold=True, color="FFFFFF", size=7)
            title.fill = PatternFill("solid", fgColor="000000")
            headers = ["Category", "Subs", "Current Charges", "Tax", "Charges + Tax"]
            for offset, header in enumerate(headers):
                cell = sheet.cell(3, start_column + offset, header)
                cell.font = Font(bold=True, color="FFFFFF", size=7)
                cell.fill = PatternFill("solid", fgColor="000000")
            sheet.cell(4, start_column, "New Sub MANUAL")
            for offset, value in enumerate(values, start=1):
                sheet.cell(4, start_column + offset, value)
            sheet.cell(5, start_column, "TOTAL")
            for offset in range(1, 5):
                source = sheet.cell(4, start_column + offset).coordinate
                sheet.cell(5, start_column + offset, f"={source}")

        sheet["N2"] = "Rotated note"
        sheet["N2"].alignment = Alignment(text_rotation=45)
        chart = BarChart()
        chart.add_data(Reference(sheet, min_col=3, min_row=3, max_row=5), titles_from_data=True)
        chart.anchor = "N8"
        sheet.add_chart(chart)
        workbook.save(path)

    def _build_hidden_workbook(self, path: Path) -> None:
        workbook = Workbook()
        cover = workbook.active
        cover.title = "Cover"
        cover["A1"] = "Visible report"
        for name, code, amount, state in (
            ("HiddenData", "A-01", 1250, "hidden"),
            ("VeryHiddenData", "B-02", 2500, "veryHidden"),
        ):
            sheet = workbook.create_sheet(name)
            sheet.append(["Code", "Amount"])
            sheet.append([code, amount])
            sheet.sheet_state = state
        workbook.save(path)

    def test_side_by_side_tables_become_separate_regions(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "billing.xlsx"
            self._build_side_by_side_workbook(path)

            survey = survey_excel_workbook(
                path,
                settings=Settings(
                    excel_small_font_points=8,
                    excel_base_dpi=300,
                    excel_max_dpi=450,
                ),
            )

            regions = survey.sheets[0].regions
            invoice_regions = [region for region in regions if region.kind == "table"]
            self.assertEqual(len(invoice_regions), 2)
            self.assertEqual(
                {region.period for region in invoice_regions},
                {"FEB 2026", "MAR 2026"},
            )
            self.assertTrue(all(region.render_dpi >= 400 for region in invoice_regions))
            self.assertTrue(
                all(region.requires_vlm_reading for region in invoice_regions)
            )
            self.assertNotIn("N2", {cell.coordinate for region in invoice_regions for cell in region.cells})
            chart_regions = [region for region in regions if region.title == "Chart 1"]
            self.assertEqual(len(chart_regions), 1)
            self.assertTrue(chart_regions[0].requires_vlm_reading)

    def test_declared_tables_are_locked_regions_even_when_touching(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "touching.xlsx"
            workbook = Workbook()
            sheet = workbook.active
            sheet.title = "Touching"
            for row in (
                ("Left ID", "Left Value", "Right ID", "Right Value"),
                (1, 10, "A", 100),
                (2, 20, "B", 200),
            ):
                sheet.append(row)
            sheet.add_table(Table(displayName="LeftTable", ref="A1:B3"))
            sheet.add_table(Table(displayName="RightTable", ref="C1:D3"))
            workbook.save(path)

            survey = survey_excel_workbook(path)
            table_regions = [
                region for region in survey.sheets[0].regions if region.kind == "table"
            ]

            self.assertEqual(
                [(region.title, region.cell_range) for region in table_regions],
                [("LeftTable", "A1:B3"), ("RightTable", "C1:D3")],
            )

    def test_diagonal_contact_does_not_merge_independent_regions(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "diagonal.xlsx"
            workbook = Workbook()
            sheet = workbook.active
            sheet["A1"] = "First"
            sheet["B1"] = "Value"
            sheet["A2"] = "A"
            sheet["B2"] = 1
            sheet["C3"] = "Second"
            sheet["D3"] = "Value"
            sheet["C4"] = "B"
            sheet["D4"] = 2
            workbook.save(path)

            survey = survey_excel_workbook(path)

            self.assertEqual(
                [region.cell_range for region in survey.sheets[0].regions],
                ["A1:B2", "C3:D4"],
            )

    def test_static_named_range_is_used_as_region_boundary(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "named-range.xlsx"
            workbook = Workbook()
            sheet = workbook.active
            sheet.title = "Named Data"
            sheet.append(["Code", "Amount", "Other", "Label", "Value"])
            sheet.append(["A", 10, "x", "One", 1])
            sheet.append(["B", 20, "y", "Two", 2])
            workbook.defined_names.add(
                DefinedName(
                    "PrimaryBlock",
                    attr_text="'Named Data'!$A$1:$B$3",
                )
            )
            workbook.save(path)

            survey = survey_excel_workbook(path)

            self.assertEqual(
                [
                    (region.title, region.cell_range)
                    for region in survey.sheets[0].regions
                ],
                [("PrimaryBlock", "A1:B3"), ("Other Label Value", "C1:E3")],
            )

    def test_header_detection_prefers_row_followed_by_table_body(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            path = root / "preamble.xlsx"
            database_path = root / "preamble.sqlite"
            workbook = Workbook()
            sheet = workbook.active
            sheet.append(["Quarterly", "Revenue", "Report", "2026"])
            sheet.append([None, None, None, None])
            sheet.append(["ID", "Area", "Amount", "Status"])
            sheet.append([1, "West", 100, "Open"])
            sheet.append([2, "East", 200, "Closed"])
            workbook.defined_names.add(
                DefinedName("ReportBlock", attr_text="'Sheet'!$A$1:$D$5")
            )
            workbook.save(path)

            survey = survey_excel_workbook(path)
            created = persist_excel_native_data(survey, database_path)

            self.assertEqual(len(created), 1)
            with closing(sqlite3.connect(database_path)) as connection:
                columns = [
                    row[1]
                    for row in connection.execute(
                        f'PRAGMA table_info("{created[0]}")'
                    )
                ]
                rows = connection.execute(
                    f'SELECT id, area, amount, status FROM "{created[0]}" '
                    "ORDER BY source_row"
                ).fetchall()
            self.assertEqual(columns[-4:], ["id", "area", "amount", "status"])
            self.assertEqual(rows, [(1, "West", 100, "Open"), (2, "East", 200, "Closed")])

    def test_styled_blank_cells_bridge_sparse_table_columns(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "styled-gap.xlsx"
            workbook = Workbook()
            sheet = workbook.active
            sheet["A1"] = "Key"
            sheet["C1"] = "Value"
            sheet["A2"] = "A"
            sheet["C2"] = 10
            thin = Side(style="thin")
            for row in range(1, 3):
                for column in range(1, 4):
                    sheet.cell(row, column).border = Border(
                        left=thin,
                        right=thin,
                        top=thin,
                        bottom=thin,
                    )
            workbook.save(path)

            survey = survey_excel_workbook(path)

            self.assertEqual(len(survey.sheets[0].regions), 1)
            self.assertEqual(survey.sheets[0].regions[0].cell_range, "A1:C2")
            self.assertEqual(survey.sheets[0].regions[0].kind, "table")

    def test_hidden_sheets_are_preserved_as_native_support(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            path = root / "hidden.xlsx"
            database_path = root / "hidden.sqlite"
            self._build_hidden_workbook(path)

            survey = survey_excel_workbook(path)
            created = persist_excel_native_data(survey, database_path)

            support_sheets = {
                sheet.name: sheet
                for sheet in survey.sheets
                if sheet.name in {"HiddenData", "VeryHiddenData"}
            }
            self.assertEqual(set(support_sheets), {"HiddenData", "VeryHiddenData"})
            self.assertTrue(all(not sheet.visible for sheet in support_sheets.values()))
            self.assertTrue(
                all(sheet.role == "support" for sheet in support_sheets.values())
            )
            self.assertTrue(
                all(len(sheet.regions) == 1 for sheet in support_sheets.values())
            )
            self.assertTrue(set(support_sheets).isdisjoint(survey.extraction_order))
            markdown, _ = compose_excel_markdown(
                survey,
                fallback_title="Hidden workbook",
            )
            self.assertNotIn("HiddenData", markdown)
            self.assertNotIn("VeryHiddenData", markdown)
            self.assertEqual(len(created), 2)
            with closing(sqlite3.connect(database_path)) as connection:
                rows = {
                    connection.execute(
                        f'SELECT sheet_name, code, amount FROM "{table_name}"'
                    ).fetchone()
                    for table_name in created
                }
            self.assertEqual(
                rows,
                {
                    ("HiddenData", "A-01", 1250),
                    ("VeryHiddenData", "B-02", 2500),
                },
            )

    def test_native_tables_group_matching_monthly_regions(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            workbook_path = root / "billing.xlsx"
            database_path = root / "billing.sqlite"
            self._build_side_by_side_workbook(workbook_path)
            survey = survey_excel_workbook(workbook_path)

            created = persist_excel_native_data(survey, database_path)

            self.assertEqual(len(created), 1)
            connection = sqlite3.connect(database_path)
            try:
                periods = {
                    row[0]
                    for row in connection.execute(
                        f'SELECT DISTINCT period FROM "{created[0]}"'
                    ).fetchall()
                }
                source_rows = connection.execute(
                    f'SELECT COUNT(*) FROM "{created[0]}"'
                ).fetchone()[0]
                evidence_count = connection.execute(
                    "SELECT COUNT(*) FROM excel_cells"
                ).fetchone()[0]
            finally:
                connection.close()
            gc.collect()

            self.assertEqual(periods, {"FEB 2026", "MAR 2026"})
            self.assertEqual(source_rows, 4)
            self.assertGreater(evidence_count, 20)

    def test_adaptive_roles_and_visual_plan_are_content_driven(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "arbitrary-name.xlsx"
            self._build_adaptive_workbook(path)

            survey = survey_excel_workbook(path)
            roles = {sheet.name: sheet.role for sheet in survey.sheets}
            visual_regions = plan_excel_visual_regions(survey)

            self.assertEqual(roles["Visual"], "dashboard")
            self.assertEqual(roles["Rollup"], "summary")
            self.assertEqual(roles["ChartData"], "support")
            self.assertEqual(roles["Records"], "detail")
            self.assertEqual(
                survey.extraction_order,
                ["Records", "Rollup", "ChartData", "Visual"],
            )
            self.assertEqual(len(visual_regions), 1)
            self.assertEqual(visual_regions[0].sheet_name, "Visual")
            self.assertEqual(visual_regions[0].render_strategy, "visual")

    def test_native_table_infers_sqlite_types_from_excel_values(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            path = root / "adaptive.xlsx"
            database_path = root / "adaptive.sqlite"
            self._build_adaptive_workbook(path)

            survey = survey_excel_workbook(path)
            created = persist_excel_native_data(survey, database_path)
            records_table = next(
                table for table in created if table.startswith("excel_records_")
            )

            with closing(sqlite3.connect(database_path)) as connection:
                column_types = {
                    row[1]: row[2]
                    for row in connection.execute(
                        f'PRAGMA table_info("{records_table}")'
                    )
                }
                first_row = connection.execute(
                    f'SELECT id, area, amount, status FROM "{records_table}" '
                    "ORDER BY source_row LIMIT 1"
                ).fetchone()

            self.assertEqual(
                column_types,
                {
                    "source_file": "TEXT",
                    "sheet_name": "TEXT",
                    "region_id": "TEXT",
                    "period": "TEXT",
                    "source_row": "INTEGER",
                    "row_role": "TEXT",
                    "id": "INTEGER",
                    "area": "TEXT",
                    "amount": "INTEGER",
                    "status": "TEXT",
                },
            )
            self.assertEqual(first_row, (1, "Area 1", 1250, "Open"))

    def test_markdown_summarizes_charts_and_links_complete_data(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "adaptive.xlsx"
            csv_path = Path(directory) / "records.csv"
            self._build_adaptive_workbook(path)
            survey = survey_excel_workbook(path)
            chart = next(sheet for sheet in survey.sheets if sheet.name == "Visual").charts[0]
            chart.series[0].name = "Value"
            chart.series[0].categories = ["2026-01", "2026-02", "Closed"]
            chart.series[0].values = [10, 15, None]
            chart.warnings = ["Seri mencampur kategori periode dengan kategori nonperiode."]

            markdown, sections = compose_excel_markdown(
                survey,
                fallback_title="adaptive",
                artifacts=[
                    ExcelNativeArtifact(
                        sheet_name="Records",
                        table_name="excel_records_test",
                        row_count=40,
                        columns=["id", "area", "amount", "status"],
                        csv_path=str(csv_path),
                    )
                ],
            )

            self.assertIn("Operational Monitoring", markdown)
            self.assertIn("### Grafik: Movement", markdown)
            self.assertIn("| Kategori | Value |", markdown)
            self.assertIn("Data lengkap", markdown)
            self.assertIn(str(csv_path), markdown)
            self.assertIn("Sheet pendukung", markdown)
            self.assertEqual(sections[0][0], "Records")

    def test_wide_region_is_tiled_and_can_be_combined(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "wide.xlsx"
            workbook = Workbook()
            sheet = workbook.active
            sheet.title = "Wide"
            for column in range(1, 25):
                sheet.cell(1, column, f"H{column}")
                for row_number in range(2, 35):
                    sheet.cell(row_number, column, row_number * column)
            workbook.save(path)

            survey = survey_excel_workbook(path, settings=Settings(
                excel_tile_max_columns=12,
                excel_tile_max_rows=40,
            ))
            tiled = split_excel_regions_into_tiles(survey, Settings(
                excel_tile_max_columns=12,
                excel_tile_max_rows=40,
            ))
            regions = tiled.sheets[0].regions
            self.assertGreaterEqual(len(regions), 2)
            self.assertTrue(all(region.is_tile for region in regions))
            self.assertEqual({region.parent_region_id for region in regions}, {"s001_r001"})
            self.assertEqual(
                {cell.coordinate for region in regions for cell in region.cells},
                {f"{get_column_letter(column)}{row}" for column in range(1, 25) for row in range(1, 35)},
            )

    def test_tiled_table_is_rendered_as_one_markdown_table(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "tall_table.xlsx"
            workbook = Workbook()
            sheet = workbook.active
            sheet.title = "Summary"
            sheet["A1"] = "Monthly settlement"
            sheet["A2"] = "Region"
            sheet["B2"] = "Amount"
            for row in range(3, 15):
                sheet.cell(row, 1, f"Region {row - 2}")
                sheet.cell(row, 2, row * 100)
            workbook.save(path)

            settings = Settings(excel_tile_max_columns=8, excel_tile_max_rows=4)
            tiled = split_excel_regions_into_tiles(
                survey_excel_workbook(path, settings=settings), settings,
            )
            tiled = tiled.model_copy(update={
                "sheets": [tiled.sheets[0].model_copy(update={"role": "summary"})],
            })
            regions = tiled.sheets[0].regions
            self.assertGreater(len(regions), 1)

            markdown, _ = compose_excel_markdown(tiled, fallback_title="Summary")

            self.assertEqual(markdown.count("| Region | Amount |"), 1)
            for row in range(1, 13):
                self.assertEqual(markdown.count(f"| Region {row} |"), 1)

    @unittest.skipUnless(_find_libreoffice_binary(), "LibreOffice tidak tersedia")
    def test_region_renderer_produces_one_page_per_region(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            workbook_path = root / "billing.xlsx"
            pdf_path = root / "billing.pdf"
            self._build_side_by_side_workbook(workbook_path)
            survey = survey_excel_workbook(workbook_path)
            invoice_regions = [
                region
                for region in survey.sheets[0].regions
                if region.kind == "table"
            ]

            rendered = _convert_excel_regions_to_pdf(
                workbook_path,
                pdf_path,
                root / "lo_profile",
                invoice_regions,
            )

            self.assertTrue(rendered)
            self.assertEqual(pdf_page_count(pdf_path), 2)

    @unittest.skipUnless(_find_libreoffice_binary(), "LibreOffice tidak tersedia")
    def test_survey_copy_recalculates_formula_cache(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            workbook_path = root / "billing.xlsx"
            self._build_side_by_side_workbook(workbook_path)

            recalculated = _convert_workbook_to_xlsx_copy(
                workbook_path, root / "recalculated"
            )
            survey = survey_excel_workbook(recalculated)
            formula_cells = [
                cell
                for region in survey.sheets[0].regions
                for cell in region.cells
                if cell.formula
            ]

            self.assertTrue(formula_cells)
            self.assertTrue(all(cell.value is not None for cell in formula_cells))


if __name__ == "__main__":
    unittest.main()
