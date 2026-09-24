"""Regresi untuk semantik khusus Excel yang tidak ada pada jalur PDF.

Kasus diambil dari workbook dashboard multi-region: baris subtotal, blok filter,
tabel yang menempel tanpa baris kosong, rentang grafik yang kebablasan, dan sheet
sumber dashboard yang juga memuat tabel mandiri.
"""

import sqlite3
import tempfile
import unittest
from contextlib import closing
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.chart import LineChart, Reference

from app.agent_graph import AgentDocumentGraph, DocumentBatchState
from app.excel import (
    _convert_workbook_to_xlsx_copy,
    _describe_excel_native_artifacts,
    _visual_regions_need_vlm_tables,
    compose_excel_markdown,
    ensure_excel_native_artifacts,
    persist_excel_native_data,
    plan_excel_visual_regions,
    survey_excel_workbook,
)
from app.ppt import _find_libreoffice_binary
from app.tabular_db import ROW_ROLE_SQL_GUIDANCE, TabularDatabaseManager

DEPARTMENTS = ["Network", "FTTH", "TX"]


class TestExcelAggregateRows(unittest.TestCase):
    def _build(self, path: Path) -> None:
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Summary"
        sheet.append(["Status", "Department", "Count", "Cost"])
        row = 2
        for status in ("DONE", "NOTYET"):
            first = row
            for index, department in enumerate(DEPARTMENTS, start=1):
                sheet.append([status, department, index, index * 100])
                row += 1
            sheet.append(
                [
                    f"{status} Total",
                    None,
                    f"=SUM(C{first}:C{row - 1})",
                    f"=SUM(D{first}:D{row - 1})",
                ]
            )
            row += 1
        sheet.append(["Grand Total", None, 12, 1200])
        sheet.append(["Total Station", "Survey", 5, 50])
        workbook.save(path)

    def test_subtotal_rows_are_flagged_not_dropped(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            path = root / "summary.xlsx"
            database_path = root / "summary.sqlite"
            self._build(path)

            survey = survey_excel_workbook(path)
            created = persist_excel_native_data(survey, database_path)
            self.assertEqual(len(created), 1)

            with closing(sqlite3.connect(database_path)) as connection:
                roles = dict(
                    connection.execute(
                        f'SELECT status, row_role FROM "{created[0]}"'
                    ).fetchall()
                )
                data_cost = connection.execute(
                    f"SELECT SUM(cost) FROM \"{created[0]}\" WHERE row_role = 'data'"
                ).fetchone()[0]

            self.assertEqual(roles["DONE Total"], "subtotal")
            self.assertEqual(roles["NOTYET Total"], "subtotal")
            self.assertEqual(roles["Grand Total"], "grand_total")
            # Label berakhiran "Total" tanpa formula SUM dan tanpa sel label kosong
            # tetap data biasa.
            self.assertEqual(roles["Total Station"], "data")
            self.assertEqual(roles["DONE"], "data")
            self.assertEqual(data_cost, 2 * (100 + 200 + 300) + 50)

            artifacts = _describe_excel_native_artifacts(
                survey, database_path, created, []
            )
            self.assertEqual(artifacts[0].aggregate_row_count, 3)
            self.assertNotIn("row_role", artifacts[0].columns)

            markdown, _ = compose_excel_markdown(survey, fallback_title="summary")
            self.assertIn("Baris agregat (3)", markdown)

    def test_inspect_database_exposes_row_role_guidance(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            path = root / "summary.xlsx"
            database_path = root / "summary.sqlite"
            self._build(path)
            created = persist_excel_native_data(
                survey_excel_workbook(path), database_path
            )

            info = TabularDatabaseManager(database_path).inspect_database()
            details = info["tables"][created[0]]
            self.assertEqual(
                details["row_role_counts"],
                {"data": 7, "subtotal": 2, "grand_total": 1},
            )
            self.assertEqual(details["aggregation_note"], ROW_ROLE_SQL_GUIDANCE)
            # Tabel tanpa row_role tidak mendapat kunci tambahan.
            self.assertNotIn("row_role_counts", info["tables"]["excel_cells"])

    def test_legacy_table_without_row_role_is_migrated(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            path = root / "summary.xlsx"
            database_path = root / "legacy.sqlite"
            self._build(path)
            survey = survey_excel_workbook(path)
            created = persist_excel_native_data(survey, database_path)

            with closing(sqlite3.connect(database_path)) as connection:
                connection.execute(f'DROP TABLE "{created[0]}"')
                connection.execute(
                    f'CREATE TABLE "{created[0]}" (source_file TEXT NOT NULL, '
                    "sheet_name TEXT NOT NULL, region_id TEXT NOT NULL, period TEXT, "
                    "source_row INTEGER NOT NULL, status TEXT, department TEXT, "
                    "count INTEGER, cost INTEGER)"
                )
                connection.commit()

            persist_excel_native_data(survey, database_path)
            with closing(sqlite3.connect(database_path)) as connection:
                count = connection.execute(
                    f"SELECT COUNT(*) FROM \"{created[0]}\" WHERE row_role != 'data'"
                ).fetchone()[0]
            self.assertEqual(count, 3)

    def test_metadata_named_headers_do_not_collide(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            path = root / "period.xlsx"
            database_path = root / "period.sqlite"
            workbook = Workbook()
            sheet = workbook.active
            sheet.append(["Period", "Amount"])
            sheet.append(["2026-01", 10])
            sheet.append(["2026-02", 20])
            workbook.save(path)

            created = persist_excel_native_data(
                survey_excel_workbook(path), database_path
            )
            with closing(sqlite3.connect(database_path)) as connection:
                rows = connection.execute(
                    f'SELECT period_value, amount FROM "{created[0]}" ORDER BY source_row'
                ).fetchall()
            self.assertEqual(rows, [("2026-01", 10), ("2026-02", 20)])


@unittest.skipUnless(_find_libreoffice_binary(), "LibreOffice tidak tersedia")
class TestExcelLayoutSemantics(unittest.TestCase):
    """Workbook dashboard kecil yang dihitung ulang LibreOffice agar cache formula ada."""

    @classmethod
    def setUpClass(cls) -> None:
        cls._directory = tempfile.TemporaryDirectory()
        root = Path(cls._directory.name)
        source = root / "dashboard.xlsx"
        cls._build_dashboard(source)
        cls.workbook_path = _convert_workbook_to_xlsx_copy(source, root / "recalc")

    @classmethod
    def tearDownClass(cls) -> None:
        cls._directory.cleanup()

    @staticmethod
    def _build_dashboard(path: Path) -> None:
        workbook = Workbook()
        data = workbook.active
        data.title = "Data"
        data.append(["Month", "Status", "Cost"])
        for month in range(1, 7):
            for status in ("DONE", "NOTYET"):
                data.append([datetime(2026, month, 1), status, month * 100])  # noqa: DTZ001 - Excel stores timezone-naive dates.

        pivot = workbook.create_sheet("Pivot")
        pivot.append(["Month", "Count", "Cost"])
        for month in range(1, 7):
            row = month + 1
            pivot.append(
                [
                    datetime(2026, month, 1),  # noqa: DTZ001 - Excel stores timezone-naive dates.
                    f"=COUNTIF(Data!$A$2:$A$13,A{row})",
                    f"=SUMIF(Data!$A$2:$A$13,A{row},Data!$C$2:$C$13)",
                ]
            )
        # Tabel kedua menempel langsung tanpa baris kosong dan lebih lebar.
        pivot.append(["Status", "Jan", "Feb", "Mar", "Grand Total"])
        for row, status in ((9, "DONE"), (10, "NOTYET")):
            pivot.append(
                [
                    status,
                    f"=COUNTIFS(Data!$B$2:$B$13,$A{row},Data!$A$2:$A$13,$A$2)",
                    f"=COUNTIFS(Data!$B$2:$B$13,$A{row},Data!$A$2:$A$13,$A$3)",
                    f"=COUNTIFS(Data!$B$2:$B$13,$A{row},Data!$A$2:$A$13,$A$4)",
                    f"=SUM(B{row}:D{row})",
                ]
            )

        dashboard = workbook.create_sheet("Dashboard")
        dashboard["A1"] = "Division Head"
        dashboard["B1"] = "All"
        dashboard["A2"] = "Region"
        dashboard["B2"] = "All"
        dashboard["D1"] = "Month"
        dashboard["E1"] = "Cost"
        for offset in range(9):
            # Referensi sengaja kebablasan ke tabel status (baris 8-10).
            dashboard.cell(2 + offset, 4, f"=Pivot!A{2 + offset}")
            dashboard.cell(2 + offset, 5, f"=Pivot!C{2 + offset}")
        for offset in range(6):
            dashboard.cell(20 + offset, 10, f"=Pivot!B{2 + offset}")
        chart = LineChart()
        chart.add_data(
            Reference(dashboard, min_col=5, min_row=1, max_row=10),
            titles_from_data=True,
        )
        chart.set_categories(Reference(dashboard, min_col=4, min_row=2, max_row=10))
        dashboard.add_chart(chart, "G2")
        workbook.save(path)

    def _survey(self):
        return survey_excel_workbook(self.workbook_path)

    def test_stacked_tables_without_blank_row_are_split(self) -> None:
        pivot = next(sheet for sheet in self._survey().sheets if sheet.name == "Pivot")
        ranges = [region.cell_range for region in pivot.regions]
        self.assertIn("A1:C7", ranges)
        self.assertIn("A8:E10", ranges)

    def test_parameter_block_is_rendered_but_not_persisted(self) -> None:
        survey = self._survey()
        dashboard = next(sheet for sheet in survey.sheets if sheet.name == "Dashboard")
        parameter = next(
            region for region in dashboard.regions if region.cell_range == "A1:B2"
        )
        self.assertEqual(parameter.semantic_role, "parameter")
        self.assertEqual(parameter.kind, "text")
        self.assertFalse(parameter.persist_native)
        self.assertFalse(parameter.referenced_by_formula)

        markdown, _ = compose_excel_markdown(survey, fallback_title="dash")
        self.assertIn("- **Region**: All", markdown)

        with tempfile.TemporaryDirectory() as directory:
            database_path = Path(directory) / "dash.sqlite"
            created = persist_excel_native_data(survey, database_path)
            with closing(sqlite3.connect(database_path)) as connection:
                for table in created:
                    columns = {
                        row[1]
                        for row in connection.execute(f'PRAGMA table_info("{table}")')
                    }
                    self.assertNotIn("division_head", columns)

    def test_chart_range_crossing_blocks_is_warned(self) -> None:
        survey = self._survey()
        self.assertTrue(
            any("melintasi 2 blok" in warning for warning in survey.warnings),
            survey.warnings,
        )
        dashboard = next(sheet for sheet in survey.sheets if sheet.name == "Dashboard")
        trend = next(
            region
            for region in dashboard.regions
            if region.kind == "table" and region.min_row == 1
        )
        self.assertEqual(trend.cell_range, "D1:E7")

    def test_support_sheet_keeps_unconsumed_blocks(self) -> None:
        survey = self._survey()
        pivot = next(sheet for sheet in survey.sheets if sheet.name == "Pivot")
        self.assertEqual(pivot.role, "support")
        persisted = {
            region.cell_range: region.persist_native
            for region in pivot.regions
            if region.kind == "table"
        }
        # Blok bulan seluruhnya dikonsumsi dashboard; blok status hanya sebagian.
        self.assertFalse(persisted["A1:C7"])
        self.assertTrue(persisted["A8:E10"])

        markdown, _ = compose_excel_markdown(survey, fallback_title="dash")
        self.assertIn("| Status | Jan | Feb | Mar | Grand Total |", markdown)

    def test_native_dashboard_skips_vlm_table_ingestion(self) -> None:
        survey = self._survey()
        visual_regions = plan_excel_visual_regions(survey)
        self.assertTrue(visual_regions)
        self.assertFalse(_visual_regions_need_vlm_tables(survey, visual_regions))
        self.assertTrue(_visual_regions_need_vlm_tables(None, visual_regions))


class TestExcelAgentMode(unittest.TestCase):
    """Agent Mode (MCP) harus memakai data native, bukan tabel hasil baca gambar."""

    AGENT_MARKDOWN = (
        "# Summary\n\n"
        "| Status | Department | Count | Cost |\n"
        "| --- | --- | --- | --- |\n"
        "| DONE | Network | 999 | 999999 |\n"
        "| DONE Total |  | 999 | 999999 |\n"
    )

    def _state(self, root: Path, path: Path) -> DocumentBatchState:
        return {
            "resolved_path": str(path),
            "resolved_out": str(root / "out"),
            "doc_type": "excel",
            "total_items": 1,
            "current_page": 1,
            "incoming_markdown": self.AGENT_MARKDOWN,
            "ingest_transactional_tables": True,
        }

    def test_save_uses_native_tables_and_skips_agent_tables(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            path = root / "summary.xlsx"
            TestExcelAggregateRows()._build(path)

            result = AgentDocumentGraph._node_save_and_stitch(self._state(root, path))

            self.assertEqual(result["status"], "success")
            sources = {table.get("source") for table in result["tabular_tables"]}
            self.assertEqual(sources, {"excel_native"})
            self.assertEqual(result["tabular_tables"][0]["aggregate_rows"], 3)

            database_path = root / "out" / "databases" / "summary.sqlite"
            with closing(sqlite3.connect(database_path)) as connection:
                tables = [
                    row[0]
                    for row in connection.execute(
                        "SELECT name FROM sqlite_master WHERE type='table'"
                    )
                ]
                native = next(
                    name for name in tables if name.startswith("excel_summary")
                )
                total = connection.execute(
                    f"SELECT SUM(cost) FROM \"{native}\" WHERE row_role = 'data'"
                ).fetchone()[0]
            # Angka 999999 dari Markdown agent tidak pernah masuk database.
            self.assertEqual(total, 1250)
            # Jalur Markdown agent (transaction_details) tidak dijalankan untuk Excel.
            self.assertNotIn("transaction_details", tables)

            native_markdown = root / "out" / "summary.native.md"
            self.assertTrue(native_markdown.exists())
            self.assertIn("Baris agregat", native_markdown.read_text(encoding="utf-8"))
            self.assertIn(str(native_markdown), result["instruction"])

    def test_native_context_is_cached_per_file_version(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            path = root / "summary.xlsx"
            TestExcelAggregateRows()._build(path)

            first = ensure_excel_native_artifacts(path, root / "out")
            context_file = root / "out" / "regions" / "excel_native_agent.json"
            first_mtime = context_file.stat().st_mtime_ns
            second = ensure_excel_native_artifacts(path, root / "out")

            self.assertIsNotNone(first)
            self.assertEqual(first, second)
            self.assertEqual(context_file.stat().st_mtime_ns, first_mtime)
            assert first is not None
            self.assertFalse(first["vlm_tables_needed"])
            self.assertIn("Summary", first["sheet_roles"])


class TestExcelNoFalseSplits(unittest.TestCase):
    def test_section_label_rows_do_not_split_table(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "sections.xlsx"
            workbook = Workbook()
            sheet = workbook.active
            sheet.append(["Item", "Qty", "Price"])
            sheet.append(["Cable", 2, 100])
            sheet.append(["Section B", None, None])
            sheet.append(["Switch", 1, 900])
            sheet.append(["Router", "N/A", "TBD"])
            sheet.append(["Modem", 3, 300])
            workbook.save(path)

            survey = survey_excel_workbook(path)
            tables = [
                region.cell_range
                for region in survey.sheets[0].regions
                if region.kind == "table"
            ]
            self.assertEqual(tables, ["A1:C6"])


if __name__ == "__main__":
    unittest.main()
