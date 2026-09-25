"""
Pipeline Excel berbasis Vision VLM.

Workbook dirender lewat LibreOffice menjadi PDF sementara, lalu mengikuti pipeline
PDF: setiap halaman hasil cetak spreadsheet dirender menjadi gambar dan dikirim
satu per satu ke VLM.
"""

from __future__ import annotations

import hashlib
import json
import logging
import re
import shutil
import socket
import sqlite3
import subprocess
import tempfile
import time
import unicodedata
from collections import deque
from contextlib import closing
from datetime import date, datetime
from datetime import time as datetime_time
from pathlib import Path
from typing import Any

from .config import DEFAULT_DPI, Settings, get_settings
from .pdf import pdf_page_count, pdf_to_images, process_multipage_pdf
from .ppt import _find_libreoffice_binary
from .schemas import (
    ExcelCellEvidence,
    ExcelChartEvidence,
    ExcelChartSeries,
    ExcelNativeArtifact,
    ExcelRegion,
    ExcelSheetSurvey,
    ExcelWorkbookSurvey,
    ExtractedDocument,
)

logger = logging.getLogger(__name__)


def _column_letter(column: int) -> str:
    from openpyxl.utils import get_column_letter

    return get_column_letter(column)


def _find_uno_python(soffice: str | None = None) -> str | None:
    """Temukan Python sistem yang menyediakan modul UNO LibreOffice."""
    candidates = []
    if soffice:
        candidates.append(str(Path(soffice).resolve().with_name("python.exe")))
    candidates.extend(
        [
            r"C:\Program Files\LibreOffice\program\python.exe",
            r"C:\Program Files (x86)\LibreOffice\program\python.exe",
            "/usr/bin/python3",
            "/usr/lib/libreoffice/program/python",
            "/Applications/LibreOffice.app/Contents/Resources/python",
        ]
    )
    for candidate in candidates:
        if Path(candidate).is_file():
            return candidate
    return None


def _convert_excel_to_pdf_with_page_scaling(
    soffice: str,
    excel_path: Path,
    pdf_path: Path,
    profile_dir: Path,
) -> bool:
    """Ekspor Excel melalui UNO dengan semua kolom dipaskan ke satu lebar halaman."""
    uno_python = _find_uno_python(soffice)
    if uno_python is None:
        return False

    try:
        with socket.socket() as probe:
            probe.bind(("127.0.0.1", 0))
            port = probe.getsockname()[1]
    except OSError as exc:
        logger.warning("Tidak dapat membuka port lokal untuk UNO: %s", exc)
        return False

    listener = subprocess.Popen(
        [
            soffice,
            f"-env:UserInstallation={profile_dir.as_uri()}",
            "--headless",
            "--norestore",
            "--nodefault",
            "--nofirststartwizard",
            f"--accept=socket,host=127.0.0.1,port={port};urp;StarOffice.ComponentContext",
        ],
        stdout=subprocess.DEVNULL,
        stderr=subprocess.DEVNULL,
    )

    uno_script = r"""
import sys
import time
import uno
from com.sun.star.beans import PropertyValue


def prop(name, value):
    item = PropertyValue()
    item.Name = name
    item.Value = value
    return item


port = int(sys.argv[1])
source_url = uno.systemPathToFileUrl(sys.argv[2])
target_url = uno.systemPathToFileUrl(sys.argv[3])
local_context = uno.getComponentContext()
resolver = local_context.ServiceManager.createInstanceWithContext(
    "com.sun.star.bridge.UnoUrlResolver", local_context
)
context = None
for _ in range(60):
    try:
        context = resolver.resolve(
            f"uno:socket,host=127.0.0.1,port={port};urp;StarOffice.ComponentContext"
        )
        break
    except Exception:
        time.sleep(0.25)
if context is None:
    raise RuntimeError("Tidak dapat terhubung ke listener LibreOffice UNO")

desktop = context.ServiceManager.createInstanceWithContext(
    "com.sun.star.frame.Desktop", context
)
document = desktop.loadComponentFromURL(
    source_url,
    "_blank",
    0,
    (prop("Hidden", True), prop("ReadOnly", True)),
)
if document is None:
    raise RuntimeError("LibreOffice tidak dapat membuka workbook")

try:
    page_styles = document.getStyleFamilies().getByName("PageStyles")
    sheets = document.getSheets()
    for index in range(sheets.getCount()):
        sheet = sheets.getByIndex(index)
        page_style = page_styles.getByName(sheet.getPropertyValue("PageStyle"))
        page_style.setPropertyValue("ScaleToPagesX", 1)
        page_style.setPropertyValue("ScaleToPagesY", 0)

    document.storeToURL(
        target_url,
        (prop("FilterName", "calc_pdf_Export"), prop("Overwrite", True)),
    )
finally:
    document.close(True)
    desktop.terminate()
"""

    try:
        subprocess.run(
            [uno_python, "-c", uno_script, str(port), str(excel_path), str(pdf_path)],
            capture_output=True,
            text=True,
            timeout=180,
            check=True,
        )
        return pdf_path.exists()
    except (OSError, subprocess.CalledProcessError, subprocess.TimeoutExpired) as exc:
        logger.warning(
            "Ekspor Excel dengan pengaturan fit-to-width gagal; gunakan ekspor default: %s",
            exc,
        )
        return False
    finally:
        listener.terminate()
        try:
            listener.wait(timeout=10)
        except subprocess.TimeoutExpired:
            listener.kill()
            listener.wait()
        time.sleep(0.5)


def convert_excel_to_pdf(
    excel_path: str | Path,
    output_dir: str | Path | None = None,
) -> Path:
    """Konversi workbook Excel/ODS menjadi PDF memakai LibreOffice headless."""
    path_obj = Path(excel_path).resolve()
    if not path_obj.exists():
        raise FileNotFoundError(f"File Excel tidak ditemukan: {path_obj}")

    soffice = _find_libreoffice_binary()
    if not soffice:
        raise RuntimeError(
            "LibreOffice (soffice) tidak ditemukan. Diperlukan untuk merender Excel ke gambar halaman."
        )

    target_dir = (
        Path(output_dir) if output_dir else Path(tempfile.mkdtemp(prefix="excel_pdf_"))
    )
    target_dir.mkdir(parents=True, exist_ok=True)
    pdf_path = target_dir / f"{path_obj.stem}.pdf"
    user_profile = target_dir / "lo_profile"
    converted = _convert_excel_to_pdf_with_page_scaling(
        soffice, path_obj, pdf_path, user_profile
    )

    if not converted:
        cmd = [
            soffice,
            f"-env:UserInstallation={user_profile.as_uri()}",
            "--headless",
            "--convert-to",
            "pdf",
            "--outdir",
            str(target_dir),
            str(path_obj),
        ]
        res = subprocess.run(
            cmd,
            capture_output=True,
            text=True,
            timeout=180,
            check=False,
        )
        if res.returncode != 0:
            raise RuntimeError(
                f"LibreOffice gagal mengonversi Excel ke PDF: {res.stderr.strip()}"
            )

    if not pdf_path.exists():
        candidates = sorted(target_dir.glob("*.pdf"))
        if candidates:
            return candidates[0]
        raise RuntimeError(
            "LibreOffice selesai, tetapi file PDF hasil konversi tidak ditemukan."
        )
    return pdf_path


def _slug(value: str, *, fallback: str = "sheet") -> str:
    normalized = unicodedata.normalize("NFKD", value)
    ascii_value = normalized.encode("ascii", "ignore").decode("ascii")
    cleaned = re.sub(r"[^a-zA-Z0-9]+", "_", ascii_value).strip("_").lower()
    return cleaned[:48] or fallback


def _json_safe_cell_value(value: Any) -> Any:
    if isinstance(value, (datetime, date, datetime_time)):
        return value.isoformat()
    if isinstance(value, (str, int, float, bool)) or value is None:
        return value
    return str(value)


def _extract_period(text: str) -> str | None:
    months = (
        "jan(?:uari|uary)?|feb(?:ruari|ruary)?|mar(?:et|ch)?|apr(?:il)?|"
        "mei|may|jun(?:i|e)?|jul(?:i|y)?|agu(?:stus)?|aug(?:ust)?|"
        "sep(?:tember)?|okt(?:ober)?|oct(?:ober)?|nov(?:ember)?|des(?:ember)?|dec(?:ember)?"
    )
    patterns = (
        rf"\b(?:{months})\s+20\d{{2}}\b",
        rf"\b\d{{1,2}}\s+(?:{months})\s+20\d{{2}}\b",
        r"\b(?:0?[1-9]|1[0-2])[/\-]20\d{2}\b",
        r"\b20\d{2}[/\-](?:0?[1-9]|1[0-2])\b",
    )
    for pattern in patterns:
        match = re.search(pattern, text, flags=re.IGNORECASE)
        if match:
            return match.group(0).strip()
    return None


def _cell_has_visible_style(cell: Any) -> bool:
    fill = getattr(cell, "fill", None)
    if fill is not None and getattr(fill, "fill_type", None):
        return True
    border = getattr(cell, "border", None)
    if border is not None:
        for side_name in ("left", "right", "top", "bottom"):
            if getattr(getattr(border, side_name, None), "style", None):
                return True
    font = getattr(cell, "font", None)
    alignment = getattr(cell, "alignment", None)
    return bool(
        getattr(font, "bold", False)
        or getattr(font, "italic", False)
        or int(getattr(alignment, "text_rotation", 0) or 0)
    )


def _connected_components(cells: set[tuple[int, int]]) -> list[set[tuple[int, int]]]:
    """Bentuk seed region tanpa menggabungkan sel yang hanya bersentuhan diagonal."""
    remaining = set(cells)
    components: list[set[tuple[int, int]]] = []
    while remaining:
        start = remaining.pop()
        component = {start}
        queue: deque[tuple[int, int]] = deque([start])
        while queue:
            row, column = queue.popleft()
            for row_offset, column_offset in ((-1, 0), (1, 0), (0, -1), (0, 1)):
                candidate = (row + row_offset, column + column_offset)
                if candidate in remaining:
                    remaining.remove(candidate)
                    component.add(candidate)
                    queue.append(candidate)
        components.append(component)
    return components


def _range_bounds(reference: str | None) -> tuple[int, int, int, int] | None:
    """Ubah referensi A1 statis menjadi batas baris/kolom yang tervalidasi."""
    if not reference:
        return None
    try:
        from openpyxl.utils.cell import range_boundaries

        min_column, min_row, max_column, max_row = range_boundaries(
            reference.replace("$", "")
        )
    except (ImportError, TypeError, ValueError):
        return None
    if min_row < 1 or min_column < 1 or max_row < min_row or max_column < min_column:
        return None
    return min_row, max_row, min_column, max_column


def _bounds_overlap(
    left: tuple[int, int, int, int],
    right: tuple[int, int, int, int],
) -> bool:
    return not (
        left[1] < right[0]
        or right[1] < left[0]
        or left[3] < right[2]
        or right[3] < left[2]
    )


def _coordinate_in_bounds(
    coordinate: tuple[int, int],
    bounds: tuple[int, int, int, int],
) -> bool:
    row, column = coordinate
    min_row, max_row, min_column, max_column = bounds
    return min_row <= row <= max_row and min_column <= column <= max_column


def _add_declared_region_spec(
    specs: list[tuple[str, str, tuple[int, int, int, int]]],
    *,
    kind: str,
    name: str,
    reference: str | None,
    max_cells: int | None = None,
) -> None:
    bounds = _range_bounds(reference)
    if bounds is None:
        return
    min_row, max_row, min_column, max_column = bounds
    area = (max_row - min_row + 1) * (max_column - min_column + 1)
    if max_cells is not None and area > max_cells:
        return
    if any(
        existing_bounds == bounds or _bounds_overlap(bounds, existing_bounds)
        for _, _, existing_bounds in specs
    ):
        return
    specs.append((kind, name, bounds))


def _static_defined_ranges(
    workbook: Any,
    sheet_title: str,
) -> list[tuple[str, str]]:
    ranges: list[tuple[str, str]] = []
    for defined_name in workbook.defined_names.values():
        name = str(getattr(defined_name, "name", "") or "Named Range")
        if name.startswith("_xlnm."):
            continue
        try:
            destinations = list(defined_name.destinations)
        except (AttributeError, TypeError, ValueError):
            continue
        for sheet_name, reference in destinations:
            normalized_sheet = str(sheet_name).strip("'").replace("''", "'")
            if normalized_sheet == sheet_title:
                ranges.append((name, reference))
    return ranges


def _declared_region_specs(
    worksheet: Any,
    workbook: Any,
) -> list[tuple[str, str, tuple[int, int, int, int]]]:
    """Ambil batas native berkepercayaan tinggi sebelum inferensi layout."""
    specs: list[tuple[str, str, tuple[int, int, int, int]]] = []

    for table in worksheet.tables.values():
        name = str(
            getattr(table, "displayName", None)
            or getattr(table, "name", None)
            or "Excel Table"
        )
        _add_declared_region_spec(
            specs,
            kind="excel_table",
            name=name,
            reference=getattr(table, "ref", None),
        )

    for name, reference in _static_defined_ranges(workbook, worksheet.title):
        _add_declared_region_spec(
            specs,
            kind="defined_name",
            name=name,
            reference=reference,
            max_cells=1_000_000,
        )

    auto_filter_ref = getattr(getattr(worksheet, "auto_filter", None), "ref", None)
    auto_filter_bounds = _range_bounds(auto_filter_ref)
    if auto_filter_bounds is not None:
        min_row, max_row, min_column, max_column = auto_filter_bounds
        if max_row > min_row and max_column > min_column:
            _add_declared_region_spec(
                specs,
                kind="auto_filter",
                name="AutoFilter",
                reference=auto_filter_ref,
                max_cells=1_000_000,
            )

    return specs


def _structural_style_coordinates(
    raw_cells: list[Any],
    data_coordinates: set[tuple[int, int]],
    declared_bounds: list[tuple[int, int, int, int]],
) -> set[tuple[int, int]]:
    """Gunakan blank styled cell hanya jika menguatkan struktur yang sudah ada."""
    structural: set[tuple[int, int]] = set()
    for cell in raw_cells:
        coordinate = (cell.row, cell.column)
        if coordinate in data_coordinates or not _cell_has_visible_style(cell):
            continue
        in_declared_region = any(
            _coordinate_in_bounds(coordinate, bounds) for bounds in declared_bounds
        )
        bridges_horizontal = (
            (cell.row, cell.column - 1) in data_coordinates
            and (cell.row, cell.column + 1) in data_coordinates
        )
        bridges_vertical = (
            (cell.row - 1, cell.column) in data_coordinates
            and (cell.row + 1, cell.column) in data_coordinates
        )
        if in_declared_region or bridges_horizontal or bridges_vertical:
            structural.add(coordinate)
    return structural


def _drawing_regions(worksheet: Any) -> list[tuple[str, int, int, int, int]]:
    """Ambil batas cell chart/gambar; gunakan ukuran aman untuk anchor satu sel."""
    drawings: list[tuple[str, int, int, int, int]] = []
    objects = [
        *(('chart', item) for item in getattr(worksheet, "_charts", [])),
        *(('image', item) for item in getattr(worksheet, "_images", [])),
    ]
    for kind, drawing in objects:
        anchor = getattr(drawing, "anchor", None)
        if anchor is None:
            continue
        start = getattr(anchor, "_from", None)
        end = getattr(anchor, "to", None)
        if start is not None:
            min_row = int(start.row) + 1
            min_column = int(start.col) + 1
            max_row = int(end.row) + 1 if end is not None else min_row + 14
            max_column = int(end.col) + 1 if end is not None else min_column + 7
        elif isinstance(anchor, str):
            try:
                from openpyxl.utils.cell import coordinate_to_tuple

                min_row, min_column = coordinate_to_tuple(anchor)
            except (TypeError, ValueError):
                continue
            max_row = min_row + 14
            max_column = min_column + 7
        else:
            continue
        drawings.append((kind, min_row, max_row, min_column, max_column))
    return drawings


def _formula_dependencies(worksheet: Any, sheet_names: set[str]) -> list[str]:
    dependencies: set[str] = set()
    pattern = re.compile(r"(?:'((?:[^']|'')+)'|([A-Za-z_][A-Za-z0-9_. -]*))!")
    for cell in getattr(worksheet, "_cells", {}).values():
        if getattr(cell, "data_type", None) != "f":
            continue
        for quoted, plain in pattern.findall(str(cell.value)):
            candidate = (quoted or plain).replace("''", "'").strip()
            if candidate in sheet_names and candidate != worksheet.title:
                dependencies.add(candidate)
    return sorted(dependencies)


# ---------------------------------------------------------------------------
# Semantik khusus spreadsheet
#
# Berbeda dengan PDF, workbook membawa bukti struktural yang tidak terlihat di
# halaman cetak: formula, rujukan antarsheet, dan sumber grafik. Helper di bawah
# memakai bukti tersebut untuk membedakan baris agregat, blok parameter/filter,
# dan tabel yang saling menempel tanpa baris kosong.
# ---------------------------------------------------------------------------

_MAX_EXCEL_ROWS = 1_048_576
_MAX_EXCEL_COLUMNS = 16_384

_FORMULA_STRING_LITERAL = re.compile(r'"(?:[^"]|"")*"')
_A1_REFERENCE_PATTERN = re.compile(
    r"(?<![A-Za-z0-9_.])"
    r"(?:(?:'(?P<quoted>(?:[^']|'')+)'|(?P<plain>[A-Za-z_][A-Za-z0-9_.]*))!)?"
    r"\$?(?P<c1>[A-Za-z]{1,3})\$?(?P<r1>\d+)"
    r"(?::\$?(?P<c2>[A-Za-z]{1,3})\$?(?P<r2>\d+))?"
    r"(?![A-Za-z0-9_(])"
)
_AGGREGATE_FORMULA_PATTERN = re.compile(
    r"^=\s*(?:SUM|SUBTOTAL|AGGREGATE)\s*\("
    r"(?:\s*\d+\s*,){0,2}"
    r"\s*\$?(?P<c1>[A-Za-z]{1,3})\$?(?P<r1>\d+)\s*:\s*\$?(?P<c2>[A-Za-z]{1,3})\$?(?P<r2>\d+)"
    r"\s*\)\s*$",
    flags=re.IGNORECASE,
)
_EXACT_AGGREGATE_LABEL = re.compile(
    r"^(?:grand\s*total|sub\s*-?\s*total|total(?:\s+(?:keseluruhan|akhir|umum))?|"
    r"jumlah(?:\s+(?:total|keseluruhan|akhir))?)\s*:?$",
    flags=re.IGNORECASE,
)
_SUFFIX_AGGREGATE_LABEL = re.compile(
    r"^.+\s+(?:sub\s*-?\s*)?total\s*:?$", flags=re.IGNORECASE
)
_GRAND_AGGREGATE_LABEL = re.compile(r"grand|keseluruhan|akhir|umum", flags=re.IGNORECASE)
_FILTER_VALUE_TOKENS = frozenset(
    {
        "all",
        "(all)",
        "semua",
        "(semua)",
        "(multiple items)",
        "(beberapa item)",
        "(blank)",
        "(kosong)",
    }
)

Bounds = tuple[int, int, int, int]
FormulaReference = tuple[str, tuple[int, int] | None, str, Bounds]
"""(sheet sumber, koordinat sumber atau None untuk grafik, sheet target, batas target)."""


def _column_index(letters: str) -> int:
    index = 0
    for character in letters.upper():
        index = index * 26 + (ord(character) - ord("A") + 1)
    return index


def _reference_match_bounds(match: re.Match[str]) -> Bounds | None:
    first_column = _column_index(match.group("c1"))
    first_row = int(match.group("r1"))
    last_column = _column_index(match.group("c2") or match.group("c1"))
    last_row = int(match.group("r2") or match.group("r1"))
    min_row, max_row = sorted((first_row, last_row))
    min_column, max_column = sorted((first_column, last_column))
    if (
        min_row < 1
        or max_row > _MAX_EXCEL_ROWS
        or min_column < 1
        or max_column > _MAX_EXCEL_COLUMNS
    ):
        return None
    return min_row, max_row, min_column, max_column


def _formula_references(
    formula: str,
    *,
    source_sheet: str,
    sheet_names: set[str],
) -> list[tuple[str, Bounds]]:
    """Ekstrak rujukan A1 statis dari satu formula (string literal diabaikan)."""
    cleaned = _FORMULA_STRING_LITERAL.sub('""', formula)
    references: list[tuple[str, Bounds]] = []
    for match in _A1_REFERENCE_PATTERN.finditer(cleaned):
        sheet_token = match.group("quoted")
        if sheet_token is not None:
            sheet_token = sheet_token.replace("''", "'")
        else:
            sheet_token = match.group("plain")
        target_sheet = sheet_token.strip() if sheet_token else source_sheet
        if target_sheet not in sheet_names:
            continue
        bounds = _reference_match_bounds(match)
        if bounds is not None:
            references.append((target_sheet, bounds))
    return references


def _workbook_formula_references(workbook: Any) -> list[FormulaReference]:
    sheet_names = set(workbook.sheetnames)
    references: list[FormulaReference] = []
    for worksheet in workbook.worksheets:
        for cell in getattr(worksheet, "_cells", {}).values():
            if getattr(cell, "data_type", None) != "f":
                continue
            for target_sheet, bounds in _formula_references(
                str(cell.value),
                source_sheet=worksheet.title,
                sheet_names=sheet_names,
            ):
                references.append(
                    (worksheet.title, (cell.row, cell.column), target_sheet, bounds)
                )
    return references


def _chart_references(
    sheet_title: str,
    charts: list[ExcelChartEvidence],
    sheet_names: set[str],
) -> list[FormulaReference]:
    references: list[FormulaReference] = []
    for chart in charts:
        for series in chart.series:
            for reference in (series.category_reference, series.value_reference):
                if not reference:
                    continue
                for target_sheet, bounds in _formula_references(
                    f"={reference}",
                    source_sheet=sheet_title,
                    sheet_names=sheet_names,
                ):
                    references.append((sheet_title, None, target_sheet, bounds))
    return references


def _region_bounds(region: ExcelRegion) -> Bounds:
    return region.min_row, region.max_row, region.min_column, region.max_column


def _is_external_reference(reference: FormulaReference, region: ExcelRegion) -> bool:
    source_sheet, source_coordinate, _, _ = reference
    if source_coordinate is None or source_sheet != region.sheet_name:
        return True
    return not _coordinate_in_bounds(source_coordinate, _region_bounds(region))


def _reference_coverage(
    region: ExcelRegion,
    references: list[FormulaReference],
) -> float:
    populated = [
        (cell.row, cell.column)
        for cell in region.cells
        if cell.value not in (None, "") or cell.formula
    ]
    if not populated or not references:
        return 0.0
    covered = sum(
        any(_coordinate_in_bounds(coordinate, bounds) for *_, bounds in references)
        for coordinate in populated
    )
    return covered / len(populated)


def _value_family(value: Any) -> str:
    if value is None or (isinstance(value, str) and not value.strip()):
        return "empty"
    if isinstance(value, bool):
        return "bool"
    if isinstance(value, (datetime, date, datetime_time)):
        return "date"
    if isinstance(value, (int, float)):
        return "number"
    return "text"


def _dominant_family(values: list[Any]) -> str | None:
    counts: dict[str, int] = {}
    for value in values:
        family = _value_family(value)
        if family != "empty":
            counts[family] = counts.get(family, 0) + 1
    if not counts:
        return None
    return max(counts.items(), key=lambda item: item[1])[0]


def _split_component_at_header_breaks(
    component: set[tuple[int, int]],
    value_at: Any,
) -> list[set[tuple[int, int]]]:
    """Pisahkan tabel yang menempel vertikal tanpa baris kosong.

    Baris dipotong bila (1) seluruh isinya teks, (2) diapit baris body yang
    memuat angka/tanggal, dan (3) blok di bawahnya memiliki cakupan kolom atau
    keluarga tipe kolom pertama yang berbeda dari blok di atasnya.
    """
    values_by_row: dict[int, list[tuple[int, Any]]] = {}
    for row, column in component:
        value = value_at((row, column))
        if _value_family(value) != "empty":
            values_by_row.setdefault(row, []).append((column, value))
    ordered_rows = sorted(values_by_row)
    if len(ordered_rows) < 4:
        return [component]

    def all_text(row: int) -> bool:
        return all(_value_family(value) == "text" for _, value in values_by_row[row])

    first_column = min(column for _, column in component)
    break_rows: list[int] = []
    band_start = 0
    for index in range(1, len(ordered_rows) - 1):
        row = ordered_rows[index]
        if len(values_by_row[row]) < 2 or not all_text(row):
            continue
        if all_text(ordered_rows[index - 1]) or all_text(ordered_rows[index + 1]):
            continue
        above = ordered_rows[band_start:index]
        below = ordered_rows[index:]
        above_columns = {column for item in above for column, _ in values_by_row[item]}
        below_columns = {column for item in below for column, _ in values_by_row[item]}
        above_family = _dominant_family(
            [
                value
                for item in above
                if not all_text(item)
                for column, value in values_by_row[item]
                if column == first_column
            ]
        )
        below_family = _dominant_family(
            [
                value
                for item in ordered_rows[index + 1 :]
                if not all_text(item)
                for column, value in values_by_row[item]
                if column == first_column
            ]
        )
        family_changed = (
            above_family is not None
            and below_family is not None
            and above_family != below_family
        )
        if above_columns != below_columns or family_changed:
            break_rows.append(row)
            band_start = index
    if not break_rows:
        return [component]

    boundaries = [*break_rows, _MAX_EXCEL_ROWS + 1]
    bands: list[set[tuple[int, int]]] = []
    lower = 0
    for upper in boundaries:
        band = {
            coordinate for coordinate in component if lower <= coordinate[0] < upper
        }
        if band:
            bands.extend(_connected_components(band))
        lower = upper
    return bands


def _is_parameter_block(
    evidence: list[ExcelCellEvidence],
    *,
    min_column: int,
    max_column: int,
) -> bool:
    """Kenali blok label-nilai kecil seperti filter laporan (mis. 'Region | All')."""
    if max_column - min_column != 1 or any(cell.formula for cell in evidence):
        return False
    rows: dict[int, dict[int, Any]] = {}
    for cell in evidence:
        if cell.value not in (None, ""):
            rows.setdefault(cell.row, {})[cell.column] = cell.value
    if not 1 <= len(rows) <= 4:
        return False
    every_label_has_colon = True
    has_filter_token = False
    for values in rows.values():
        label = values.get(min_column)
        value = values.get(max_column)
        if not isinstance(label, str) or not label.strip() or value in (None, ""):
            return False
        every_label_has_colon = every_label_has_colon and label.strip().endswith(":")
        if isinstance(value, str) and value.strip().casefold() in _FILTER_VALUE_TOKENS:
            has_filter_token = True
    return has_filter_token or every_label_has_colon


def _aggregate_row_roles(
    region: ExcelRegion,
    header_row: int,
) -> dict[int, str]:
    """Tandai baris subtotal/grand total agar tidak terhitung ulang sebagai data."""
    by_row: dict[int, list[ExcelCellEvidence]] = {}
    for cell in region.cells:
        if cell.row > header_row:
            by_row.setdefault(cell.row, []).append(cell)
    data_rows = sorted(
        row
        for row, cells in by_row.items()
        if any(cell.value not in (None, "") for cell in cells)
    )
    if not data_rows:
        return {}

    roles: dict[int, str] = {}
    previous_text_columns: set[int] = set()
    for row in data_rows:
        cells = sorted(by_row[row], key=lambda item: item.column)
        text_columns = {
            cell.column
            for cell in cells
            if isinstance(cell.value, str) and cell.value.strip()
        }
        # Formula tanpa cache (mis. workbook hasil skrip) tetap dihitung sebagai
        # nilai agregat potensial.
        has_number = any(
            (isinstance(cell.value, (int, float)) and not isinstance(cell.value, bool))
            or bool(cell.formula)
            for cell in cells
        )
        label = next(
            (
                str(cell.value).strip()
                for cell in cells
                if isinstance(cell.value, str) and cell.value.strip()
            ),
            None,
        )
        sum_hit = False
        sum_from_first_row = False
        for cell in cells:
            match = _AGGREGATE_FORMULA_PATTERN.match(cell.formula or "")
            if match is None:
                continue
            start_row, end_row = int(match.group("r1")), int(match.group("r2"))
            if (
                _column_index(match.group("c1")) == cell.column
                and _column_index(match.group("c2")) == cell.column
                and header_row < start_row <= end_row < row
            ):
                sum_hit = True
                sum_from_first_row = sum_from_first_row or start_row == data_rows[0]
        exact_label = bool(label and _EXACT_AGGREGATE_LABEL.fullmatch(label))
        suffix_label = bool(
            label
            and _SUFFIX_AGGREGATE_LABEL.fullmatch(label)
            and (sum_hit or bool(previous_text_columns - text_columns))
        )
        if has_number and (exact_label or suffix_label or sum_hit):
            is_last = row == data_rows[-1]
            grand = bool(
                label
                and (exact_label or suffix_label)
                and _GRAND_AGGREGATE_LABEL.search(label)
            ) or (
                is_last
                and (
                    (exact_label and not re.match(r"^sub", label or "", re.IGNORECASE))
                    or (sum_from_first_row and bool(roles))
                )
            )
            roles[row] = "grand_total" if grand else "subtotal"
        else:
            previous_text_columns = text_columns
    return roles


def _annotate_chart_source_boundaries(
    sheets: list[ExcelSheetSurvey],
    warnings: list[str],
) -> None:
    """Peringatkan bila rentang sumber grafik melintasi lebih dari satu blok."""
    sheets_by_name = {sheet.name: sheet for sheet in sheets}
    for sheet in sheets:
        for chart in sheet.charts:
            for series in chart.series:
                for reference in (series.category_reference, series.value_reference):
                    if not reference:
                        continue
                    for target_sheet, bounds in _formula_references(
                        f"={reference}",
                        source_sheet=sheet.name,
                        sheet_names=set(sheets_by_name),
                    ):
                        overlapping = [
                            region
                            for region in sheets_by_name[target_sheet].regions
                            if region.kind != "mixed"
                            and _bounds_overlap(_region_bounds(region), bounds)
                        ]
                        if len(overlapping) < 2:
                            continue
                        message = (
                            f"Rentang sumber {reference} melintasi {len(overlapping)} "
                            "blok berbeda ("
                            + ", ".join(region.cell_range for region in overlapping)
                            + "); titik di luar blok pertama kemungkinan bukan bagian seri."
                        )
                        if message not in chart.warnings:
                            chart.warnings.append(message)
                            warnings.append(f"{sheet.name} - {chart.title}: {message}")


def _chart_title(chart: Any, fallback: str) -> str:
    try:
        paragraphs = chart.title.tx.rich.p
        value = "".join(
            run.t or "" for paragraph in paragraphs for run in paragraph.r
        ).strip()
        return value or fallback
    except (AttributeError, TypeError):
        return fallback


def _cache_values(reference: Any) -> tuple[str | None, list[Any]]:
    if reference is None:
        return None, []
    formula = getattr(reference, "f", None)
    cache = getattr(reference, "strCache", None) or getattr(
        reference, "numCache", None
    )
    points = getattr(cache, "pt", None) or []
    if not points:
        return formula, []
    point_count = int(getattr(cache, "ptCount", 0) or 0)
    max_index = max(int(point.idx) for point in points)
    values: list[Any] = [None] * max(point_count, max_index + 1)
    for point in points:
        value: Any = point.v
        if isinstance(value, float) and value.is_integer():
            value = int(value)
        values[int(point.idx)] = value
    return formula, values


def _reference_values(formula: str | None, workbook: Any) -> list[Any]:
    if not formula or "!" not in formula:
        return []
    sheet_token, cell_range = formula.rsplit("!", 1)
    sheet_name = sheet_token.strip("'").replace("''", "'")
    if sheet_name not in workbook.sheetnames:
        return []
    try:
        cells = workbook[sheet_name][cell_range.replace("$", "")]
    except (KeyError, TypeError, ValueError):
        return []
    if not isinstance(cells, tuple):
        return [cells.value]
    values: list[Any] = []
    for row in cells:
        if isinstance(row, tuple):
            values.extend(cell.value for cell in row)
        else:
            values.append(row.value)
    return values


def _series_title(series: Any, index: int, workbook: Any = None) -> str:
    text = getattr(series, "tx", None)
    reference = getattr(text, "strRef", None)
    formula, cached = _cache_values(reference)
    if cached and cached[0] not in (None, ""):
        return str(cached[0])
    referenced = _reference_values(formula, workbook) if workbook is not None else []
    if referenced and referenced[0] not in (None, ""):
        return str(referenced[0])
    literal = getattr(text, "v", None)
    return str(literal).strip() if literal else f"Series {index}"


def _category_family(value: Any) -> str:
    if isinstance(value, (datetime, date)):
        return "date"
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        return "number"
    text = str(value or "").strip()
    if re.fullmatch(
        r"(?:jan|feb|mar|apr|may|mei|jun|jul|aug|agu|sep|oct|okt|nov|dec|des)[a-z]*[- /]\d{2,4}",
        text,
        flags=re.IGNORECASE,
    ):
        return "date"
    return "text"


def _extract_chart_evidence(
    worksheet: Any,
    *,
    value_workbook: Any = None,
) -> list[ExcelChartEvidence]:
    try:
        from openpyxl.utils import get_column_letter
    except ImportError as exc:  # pragma: no cover
        raise RuntimeError("openpyxl diperlukan untuk membaca grafik workbook.") from exc

    charts: list[ExcelChartEvidence] = []
    for chart_index, chart in enumerate(getattr(worksheet, "_charts", []), start=1):
        anchor = getattr(chart, "anchor", None)
        start = getattr(anchor, "_from", None)
        end = getattr(anchor, "to", None)
        if start is not None:
            min_row = int(start.row) + 1
            min_column = int(start.col) + 1
            max_row = int(end.row) + 1 if end is not None else min_row + 14
            max_column = int(end.col) + 1 if end is not None else min_column + 7
        else:
            min_row = min_column = 1
            max_row, max_column = 15, 8
        cell_range = (
            f"{get_column_letter(min_column)}{min_row}:"
            f"{get_column_letter(max_column)}{max_row}"
        )
        evidence_series: list[ExcelChartSeries] = []
        warnings: list[str] = []
        for series_index, series in enumerate(chart.series, start=1):
            category_source = getattr(series, "cat", None)
            category_reference = getattr(category_source, "strRef", None) or getattr(
                category_source, "numRef", None
            )
            value_source = getattr(series, "val", None)
            value_reference = getattr(value_source, "numRef", None)
            category_formula, categories = _cache_values(category_reference)
            value_formula, raw_values = _cache_values(value_reference)
            if not categories and value_workbook is not None:
                categories = _reference_values(category_formula, value_workbook)
            if not raw_values and value_workbook is not None:
                raw_values = _reference_values(value_formula, value_workbook)
            values = [
                value if isinstance(value, (int, float)) else None
                for value in raw_values
            ]
            series_name = _series_title(series, series_index, value_workbook)
            if len(categories) != len(values):
                warnings.append(
                    f"Seri {series_name} memiliki jumlah kategori "
                    "dan nilai yang berbeda."
                )
            families = {
                _category_family(category)
                for category, value in zip(categories, values, strict=False)
                if category not in (None, "") and value is not None
            }
            if "date" in families and len(families) > 1:
                warnings.append(
                    f"Seri {series_name} mencampur kategori "
                    "periode dengan kategori nonperiode."
                )
            if any(value is None for value in values):
                warnings.append(f"Seri {series_name} memiliki nilai kosong.")
            evidence_series.append(
                ExcelChartSeries(
                    name=series_name,
                    category_reference=category_formula,
                    value_reference=value_formula,
                    categories=[_json_safe_cell_value(value) for value in categories],
                    values=values,
                )
            )
        charts.append(
            ExcelChartEvidence(
                chart_id=f"s{worksheet._parent.index(worksheet) + 1:03d}_c{chart_index:03d}",
                title=_chart_title(chart, f"Chart {chart_index}"),
                chart_type=type(chart).__name__.removesuffix("Chart").lower(),
                cell_range=cell_range,
                series=evidence_series,
                warnings=list(dict.fromkeys(warnings)),
            )
        )
    return charts


_SUPPORT_CONSUMED_COVERAGE = 0.8


def _classify_excel_sheets(
    sheets: list[ExcelSheetSurvey],
    *,
    references: list[FormulaReference] | None = None,
) -> list[str]:
    dashboards = {sheet.name for sheet in sheets if sheet.charts and sheet.visible}
    dashboard_dependencies = {
        dependency
        for sheet in sheets
        if sheet.name in dashboards
        for dependency in sheet.dependencies
    }
    for sheet in sheets:
        if not sheet.visible:
            continue
        formula_ratio = sheet.formula_count / max(1, sheet.nonempty_cell_count)
        largest = max(
            (region for region in sheet.regions if region.kind == "table"),
            key=lambda region: len(region.cells),
            default=None,
        )
        tall_dataset = bool(
            largest
            and largest.max_row - largest.min_row + 1
            >= max(12, 3 * (largest.max_column - largest.min_column + 1))
            and len(largest.cells)
            >= 0.55
            * (largest.max_row - largest.min_row + 1)
            * (largest.max_column - largest.min_column + 1)
        )
        if sheet.charts:
            sheet.role = "dashboard"
            sheet.role_confidence = 0.98
            sheet.role_reasons = ["memiliki grafik native"]
            sheet.render_strategy = "hybrid"
        elif (
            sheet.name in dashboard_dependencies
            and formula_ratio >= 0.35
            and sheet.dependencies
        ):
            sheet.role = "support"
            sheet.role_confidence = 0.9
            sheet.role_reasons = [
                "menjadi sumber dashboard",
                "didominasi formula agregasi",
            ]
            sheet.render_strategy = "native"
        elif formula_ratio >= 0.35 and sheet.dependencies:
            sheet.role = "summary"
            sheet.role_confidence = min(0.95, 0.65 + formula_ratio / 3)
            sheet.role_reasons = ["merangkum sheet lain melalui formula"]
            if len([region for region in sheet.regions if region.kind == "table"]) > 1:
                sheet.role_reasons.append("memiliki beberapa blok tabel ringkasan")
            sheet.render_strategy = "native"
        elif tall_dataset:
            sheet.role = "detail"
            sheet.role_confidence = 0.9
            sheet.role_reasons = [
                "memiliki header dan pola baris berulang",
                "bentuk data lebih tinggi daripada lebar",
            ]
            sheet.render_strategy = "native"
        else:
            sheet.role = "plain"
            sheet.role_confidence = 0.65
            sheet.role_reasons = ["tidak menunjukkan pola dashboard atau agregasi"]
            sheet.render_strategy = "hybrid"

        for region in sheet.regions:
            if sheet.role in {"detail", "summary", "support"} or sheet.role == "dashboard":
                region.render_strategy = "native"
            else:
                region.render_strategy = "hybrid"
            region.persist_native = region.kind == "table" and sheet.role != "support"
            if sheet.role == "support" and region.kind == "table":
                # Sheet sumber dashboard bisa memuat blok lain yang tidak pernah
                # ditampilkan. Hanya blok yang benar-benar dikonsumsi dashboard
                # (formula atau grafik) yang dianggap duplikat dan tidak disimpan.
                dashboard_references = [
                    reference
                    for reference in references or []
                    if reference[0] in dashboards
                    and reference[2] == sheet.name
                    and _bounds_overlap(reference[3], _region_bounds(region))
                ]
                region.persist_native = (
                    _reference_coverage(region, dashboard_references)
                    < _SUPPORT_CONSUMED_COVERAGE
                )

    # Hasil Markdown mengikuti urutan tab Excel; peran sheet tetap menentukan
    # strategi ekstraksi dan tampilan visual, bukan urutan baca dokumen.
    return [sheet.name for sheet in sheets if sheet.visible]


def _region_native_text(
    *,
    sheet_name: str,
    cell_range: str,
    cells: list[ExcelCellEvidence],
) -> str:
    lines = [f"Sheet: {sheet_name}", f"Range: {cell_range}"]
    by_row: dict[int, list[ExcelCellEvidence]] = {}
    for cell in cells:
        by_row.setdefault(cell.row, []).append(cell)
    for row_number in sorted(by_row):
        items = []
        for cell in sorted(by_row[row_number], key=lambda item: item.column):
            value = cell.value
            if value is None and cell.formula is None:
                continue
            rendered = "" if value is None else str(value)
            if cell.formula:
                rendered = f"{rendered} [formula={cell.formula}]"
            items.append(f"{cell.coordinate}={rendered}")
        if items:
            lines.append(" | ".join(items))
    return "\n".join(lines)


def survey_excel_workbook(
    excel_path: str | Path,
    *,
    settings: Settings | None = None,
) -> ExcelWorkbookSurvey:
    """Inventaris region workbook tanpa mengubah file sumber."""
    try:
        from openpyxl import load_workbook
        from openpyxl.utils import get_column_letter, range_boundaries
    except ImportError as exc:  # pragma: no cover - dependency installation guard
        raise RuntimeError(
            "openpyxl diperlukan untuk survei struktur workbook. Jalankan instalasi dependensi proyek."
        ) from exc

    resolved_settings = settings or get_settings()
    source = Path(excel_path).resolve()
    keep_vba = source.suffix.lower() == ".xlsm"
    formula_book = load_workbook(
        source,
        read_only=False,
        data_only=False,
        keep_vba=keep_vba,
        keep_links=True,
    )
    value_book = load_workbook(
        source,
        read_only=False,
        data_only=True,
        keep_vba=keep_vba,
        keep_links=True,
    )
    warnings: list[str] = []
    surveyed_sheets: list[ExcelSheetSurvey] = []
    workbook_references: list[FormulaReference] = []
    try:
        workbook_references = _workbook_formula_references(formula_book)
        for sheet_index, worksheet in enumerate(formula_book.worksheets):
            visible = worksheet.sheet_state == "visible"
            cached_sheet = value_book[worksheet.title]

            def value_at(
                coordinate: tuple[int, int],
                *,
                _formula_sheet: Any = worksheet,
                _value_sheet: Any = cached_sheet,
            ) -> Any:
                source_cell = getattr(_formula_sheet, "_cells", {}).get(coordinate)
                if source_cell is None:
                    return None
                if getattr(source_cell, "data_type", None) == "f":
                    cached_cell = getattr(_value_sheet, "_cells", {}).get(coordinate)
                    return getattr(cached_cell, "value", None)
                return source_cell.value

            raw_cells = list(getattr(worksheet, "_cells", {}).values())
            drawing_specs = _drawing_regions(worksheet)
            declared_specs = _declared_region_specs(worksheet, formula_book)
            chart_evidence = _extract_chart_evidence(
                worksheet, value_workbook=value_book
            )
            for chart in chart_evidence:
                warnings.extend(
                    f"{worksheet.title} - {chart.title}: {message}"
                    for message in chart.warnings
                )
            data_coordinates = {
                (cell.row, cell.column)
                for cell in raw_cells
                if cell.value is not None
            }
            if not data_coordinates and not drawing_specs and not declared_specs:
                surveyed_sheets.append(
                    ExcelSheetSurvey(
                        name=worksheet.title,
                        index=sheet_index,
                        visible=visible,
                        role="plain" if visible else "support",
                        role_confidence=0.0 if visible else 1.0,
                        role_reasons=(
                            []
                            if visible
                            else ["sheet tersembunyi dipertahankan untuk audit native"]
                        ),
                        render_strategy="hybrid" if visible else "native",
                    )
                )
                continue

            merged_by_coordinate: dict[tuple[int, int], str] = {}
            merged_bounds_by_anchor: dict[tuple[int, int], tuple[int, int, int, int]] = {}
            bounds: list[tuple[int, int, int, int]] = [
                (row, row, column, column) for row, column in data_coordinates
            ]
            bounds.extend(spec_bounds for _, _, spec_bounds in declared_specs)
            for merged_range in worksheet.merged_cells.ranges:
                min_column, min_row, max_column, max_row = range_boundaries(
                    str(merged_range)
                )
                if (max_row - min_row + 1) * (max_column - min_column + 1) > 10_000:
                    warnings.append(
                        f"Merge sangat besar di {worksheet.title}!{merged_range} diabaikan sebagian."
                    )
                    continue
                anchor_value = worksheet.cell(min_row, min_column).value
                if anchor_value is not None:
                    merged_bounds_by_anchor[(min_row, min_column)] = (
                        min_row,
                        max_row,
                        min_column,
                        max_column,
                    )
                    bounds.append((min_row, max_row, min_column, max_column))
                for row in range(min_row, max_row + 1):
                    for column in range(min_column, max_column + 1):
                        merged_by_coordinate[(row, column)] = str(merged_range)

            bounds.extend(
                (min_row, max_row, min_column, max_column)
                for _, min_row, max_row, min_column, max_column in drawing_specs
            )
            min_data_row = min(item[0] for item in bounds)
            max_data_row = max(item[1] for item in bounds)
            min_data_column = min(item[2] for item in bounds)
            max_data_column = max(item[3] for item in bounds)

            used_width = max_data_column - min_data_column + 1
            isolated_merged_anchors = {
                anchor
                for anchor, (min_row, max_row, min_column, max_column) in merged_bounds_by_anchor.items()
                if min_row == max_row
                and max_column - min_column + 1 >= 8
                and max_column - min_column + 1 >= used_width * 0.5
            }
            declared_bounds = [item[2] for item in declared_specs]
            structural_coordinates = data_coordinates | _structural_style_coordinates(
                raw_cells,
                data_coordinates,
                declared_bounds,
            )
            claimed_coordinates: set[tuple[int, int]] = set()
            region_candidates: list[
                tuple[
                    set[tuple[int, int]],
                    tuple[int, int, int, int] | None,
                    str | None,
                    str | None,
                ]
            ] = []
            for declared_kind, declared_name, declared_region_bounds in declared_specs:
                component = {
                    coordinate
                    for coordinate in structural_coordinates
                    if _coordinate_in_bounds(coordinate, declared_region_bounds)
                }
                if not component.intersection(data_coordinates):
                    continue
                region_candidates.append(
                    (
                        component,
                        declared_region_bounds,
                        declared_name,
                        declared_kind,
                    )
                )
                claimed_coordinates.update(component)

            inferred_coordinates = (
                structural_coordinates
                - claimed_coordinates
                - isolated_merged_anchors
            )
            region_candidates.extend(
                (band, None, None, None)
                for component in _connected_components(inferred_coordinates)
                for band in _split_component_at_header_breaks(component, value_at)
            )
            region_candidates.extend(
                ({anchor}, None, None, None)
                for anchor in isolated_merged_anchors
                if anchor not in claimed_coordinates
            )

            regions: list[ExcelRegion] = []
            region_candidates.sort(
                key=lambda candidate: (
                    candidate[1][0]
                    if candidate[1] is not None
                    else min(row for row, _ in candidate[0]),
                    candidate[1][2]
                    if candidate[1] is not None
                    else min(column for _, column in candidate[0]),
                ),
            )
            for component, forced_bounds, declared_name, declared_kind in region_candidates:
                if not component.intersection(data_coordinates):
                    continue
                if forced_bounds is not None:
                    min_row, max_row, min_column, max_column = forced_bounds
                else:
                    min_row = min(row for row, _ in component)
                    max_row = max(row for row, _ in component)
                    min_column = min(column for _, column in component)
                    max_column = max(column for _, column in component)
                    for anchor in component:
                        merged_bounds = merged_bounds_by_anchor.get(anchor)
                        if merged_bounds is None:
                            continue
                        (
                            merged_min_row,
                            merged_max_row,
                            merged_min_col,
                            merged_max_col,
                        ) = merged_bounds
                        min_row = min(min_row, merged_min_row)
                        max_row = max(max_row, merged_max_row)
                        min_column = min(min_column, merged_min_col)
                        max_column = max(max_column, merged_max_col)
                evidence: list[ExcelCellEvidence] = []
                text_values: list[str] = []
                minimum_font_size: float | None = None
                has_rotation = False
                formula_without_cache = False

                for row, column in sorted(component):
                    source_cell = worksheet.cell(row, column)
                    formula = (
                        str(source_cell.value)
                        if source_cell.data_type == "f"
                        else None
                    )
                    cached_value = cached_sheet.cell(row, column).value
                    value = cached_value if formula else source_cell.value
                    if formula and cached_value is None:
                        formula_without_cache = True
                    if value is None and formula is None:
                        continue
                    font_size_raw = getattr(source_cell.font, "sz", None)
                    font_size = float(font_size_raw) if font_size_raw else None
                    if font_size is not None:
                        minimum_font_size = (
                            font_size
                            if minimum_font_size is None
                            else min(minimum_font_size, font_size)
                        )
                    rotation = int(source_cell.alignment.text_rotation or 0)
                    has_rotation = has_rotation or rotation != 0
                    safe_value = _json_safe_cell_value(value)
                    if isinstance(safe_value, str) and safe_value.strip():
                        text_values.append(safe_value.strip())
                    evidence.append(
                        ExcelCellEvidence(
                            coordinate=source_cell.coordinate,
                            row=row,
                            column=column,
                            value=safe_value,
                            formula=formula,
                            number_format=source_cell.number_format or "General",
                            data_type=source_cell.data_type,
                            merged_range=merged_by_coordinate.get((row, column)),
                            font_size=font_size,
                            text_rotation=rotation,
                        )
                    )

                if not evidence:
                    continue
                cell_range = (
                    f"{get_column_letter(min_column)}{min_row}:"
                    f"{get_column_letter(max_column)}{max_row}"
                )
                first_rows = [
                    str(cell.value).strip()
                    for cell in evidence
                    if cell.value not in (None, "") and cell.row <= min_row + 2
                ]
                title = declared_name or " ".join(first_rows[:3]).strip() or None
                period = _extract_period(" ".join(text_values[:20]))
                row_count = len({cell.row for cell in evidence})
                column_count = len({cell.column for cell in evidence})
                is_parameter = declared_kind is None and _is_parameter_block(
                    evidence,
                    min_column=min_column,
                    max_column=max_column,
                )
                kind = (
                    "table"
                    if not is_parameter
                    and (
                        declared_kind in {"excel_table", "auto_filter"}
                        or (row_count >= 2 and column_count >= 2)
                    )
                    else "text"
                )
                width = max_column - min_column + 1
                height = max_row - min_row + 1
                small_font = bool(
                    minimum_font_size is not None
                    and minimum_font_size <= resolved_settings.excel_small_font_points
                )
                requires_vlm = bool(
                    small_font
                    or has_rotation
                    or width > resolved_settings.excel_max_region_columns
                    or height > resolved_settings.excel_max_region_rows
                )
                render_dpi = resolved_settings.excel_base_dpi
                if requires_vlm:
                    render_dpi = min(
                        resolved_settings.excel_max_dpi,
                        max(render_dpi, 400),
                    )
                if minimum_font_size is not None and minimum_font_size <= 6:
                    render_dpi = resolved_settings.excel_max_dpi

                region = ExcelRegion(
                    region_id=f"s{sheet_index + 1:03d}_r{len(regions) + 1:03d}",
                    sheet_name=worksheet.title,
                    sheet_index=sheet_index,
                    cell_range=cell_range,
                    min_row=min_row,
                    max_row=max_row,
                    min_column=min_column,
                    max_column=max_column,
                    title=title,
                    period=period,
                    kind=kind,
                    render_dpi=render_dpi,
                    requires_vlm_reading=requires_vlm,
                    cells=evidence,
                    semantic_role="parameter" if is_parameter else None,
                )
                region.native_text = _region_native_text(
                    sheet_name=worksheet.title,
                    cell_range=cell_range,
                    cells=evidence,
                )
                regions.append(region)
                if formula_without_cache:
                    warnings.append(
                        f"Nilai cache formula tidak tersedia di {worksheet.title}!{cell_range}."
                    )

            for kind, min_row, max_row, min_column, max_column in drawing_specs:
                drawing_number = (
                    sum(
                        1
                        for region in regions
                        if (region.title or "").startswith(f"{kind.title()} ")
                    )
                    + 1
                )
                cell_range = (
                    f"{get_column_letter(min_column)}{min_row}:"
                    f"{get_column_letter(max_column)}{max_row}"
                )
                regions.append(
                    ExcelRegion(
                        region_id=f"s{sheet_index + 1:03d}_r{len(regions) + 1:03d}",
                        sheet_name=worksheet.title,
                        sheet_index=sheet_index,
                        cell_range=cell_range,
                        min_row=min_row,
                        max_row=max_row,
                        min_column=min_column,
                        max_column=max_column,
                        title=f"{kind.title()} {drawing_number}",
                        kind="mixed",
                        render_dpi=min(
                            resolved_settings.excel_max_dpi,
                            max(resolved_settings.excel_base_dpi, 400),
                        ),
                        requires_vlm_reading=True,
                        native_text=(
                            f"Sheet: {worksheet.title}\nRange: {cell_range}\n"
                            f"Visual object: {kind}"
                        ),
                    )
                )

            used_range = (
                f"{get_column_letter(min_data_column)}{min_data_row}:"
                f"{get_column_letter(max_data_column)}{max_data_row}"
            )
            surveyed_sheets.append(
                ExcelSheetSurvey(
                    name=worksheet.title,
                    index=sheet_index,
                    visible=visible,
                    used_range=used_range,
                    role="plain" if visible else "support",
                    role_confidence=0.0 if visible else 1.0,
                    role_reasons=(
                        []
                        if visible
                        else ["sheet tersembunyi dipertahankan untuk audit native"]
                    ),
                    render_strategy="hybrid" if visible else "native",
                    nonempty_cell_count=len(data_coordinates),
                    formula_count=sum(
                        cell.data_type == "f" for cell in raw_cells
                    ),
                    row_count=max_data_row - min_data_row + 1,
                    column_count=max_data_column - min_data_column + 1,
                    dependencies=_formula_dependencies(
                        worksheet, set(formula_book.sheetnames)
                    ),
                    charts=chart_evidence,
                    regions=regions,
                )
            )
    finally:
        formula_book.close()
        value_book.close()

    sheet_names = {sheet.name for sheet in surveyed_sheets}
    all_references = [
        *workbook_references,
        *(
            reference
            for sheet in surveyed_sheets
            for reference in _chart_references(sheet.name, sheet.charts, sheet_names)
        ),
    ]
    for sheet in surveyed_sheets:
        sheet_references = [
            reference for reference in all_references if reference[2] == sheet.name
        ]
        for region in sheet.regions:
            region.referenced_by_formula = any(
                _bounds_overlap(reference[3], _region_bounds(region))
                and _is_external_reference(reference, region)
                for reference in sheet_references
            )
    _annotate_chart_source_boundaries(surveyed_sheets, warnings)
    extraction_order = _classify_excel_sheets(
        surveyed_sheets, references=all_references
    )
    return ExcelWorkbookSurvey(
        source_file=str(source),
        workbook_format=source.suffix.lower().lstrip("."),
        sheets=surveyed_sheets,
        extraction_order=extraction_order,
        warnings=list(dict.fromkeys(warnings)),
    )


def _split_region_by_dimensions(
    region: ExcelRegion, settings: Settings
) -> list[ExcelRegion]:
    """Split an oversized Excel region while retaining repeated header evidence."""
    max_columns = max(1, settings.excel_tile_max_columns)
    max_rows = max(1, settings.excel_tile_max_rows)
    width = region.max_column - region.min_column + 1
    height = region.max_row - region.min_row + 1
    if width <= max_columns and height <= max_rows:
        return [region]

    evidence_by_row: dict[int, list[ExcelCellEvidence]] = {}
    for item in region.cells:
        evidence_by_row.setdefault(item.row, []).append(item)
    populated_rows = sorted(evidence_by_row)
    header_count = min(3, len(populated_rows))
    header_rows = set(populated_rows[:header_count])
    data_rows = populated_rows[header_count:] or populated_rows
    row_ranges = []
    for offset in range(0, len(data_rows), max_rows):
        chunk = data_rows[offset : offset + max_rows]
        row_ranges.append((min(chunk), max(chunk)))
    column_ranges = [
        (start, min(start + max_columns - 1, region.max_column))
        for start in range(region.min_column, region.max_column + 1, max_columns)
    ]

    tiles: list[ExcelRegion] = []
    for tile_row_index, (data_min_row, data_max_row) in enumerate(row_ranges):
        row_set = set(header_rows)
        row_set.update(range(data_min_row, data_max_row + 1))
        for tile_column_index, (min_column, max_column) in enumerate(column_ranges):
            tile_cells = [
                item
                for item in region.cells
                if item.row in row_set and min_column <= item.column <= max_column
            ]
            if not tile_cells:
                continue
            tile_min_row = min(item.row for item in tile_cells)
            tile_max_row = max(item.row for item in tile_cells)
            tile_id = f"{region.region_id}_t{tile_row_index + 1:02d}_{tile_column_index + 1:02d}"
            tile = region.model_copy(
                update={
                    "region_id": tile_id,
                    "cell_range": (
                        f"{_column_letter(min_column)}{tile_min_row}:"
                        f"{_column_letter(max_column)}{tile_max_row}"
                    ),
                    "min_row": tile_min_row,
                    "max_row": tile_max_row,
                    "min_column": min_column,
                    "max_column": max_column,
                    "cells": tile_cells,
                    "native_text": "",
                    "parent_region_id": region.region_id,
                    "tile_row_index": tile_row_index,
                    "tile_column_index": tile_column_index,
                    "is_tile": True,
                }
            )
            tile.native_text = _region_native_text(
                sheet_name=tile.sheet_name,
                cell_range=tile.cell_range,
                cells=tile.cells,
            )
            tiles.append(tile)

    logger.info(
        "[Excel] Region %s dipecah menjadi %d tile (%d kolom x %d baris maksimum).",
        region.region_id,
        len(tiles),
        max_columns,
        max_rows,
    )
    return tiles or [region]


class ExcelTileBudgetError(Exception):
    """A single cell cannot fit the configured evidence budget."""


def _estimate_native_tokens(text: str) -> int:
    """Conservative local estimate, not the model's multimodal tokenizer."""
    return (len(text.encode("utf-8")) + 2) // 3


def _split_region_into_tiles(
    region: ExcelRegion, settings: Settings
) -> list[ExcelRegion]:
    budget = max(1, settings.excel_tile_max_native_tokens)

    def refine(candidate: ExcelRegion) -> list[ExcelRegion]:
        if _estimate_native_tokens(candidate.native_text) <= budget:
            return [candidate]
        rows = sorted({cell.row for cell in candidate.cells})
        columns = sorted({cell.column for cell in candidate.cells})
        # Retain the existing three-row header convention when possible.
        # Actual serialized evidence includes coordinates, values and formulas.
        groups: list[list[ExcelCellEvidence]]
        if len(rows) > 4:
            headers = set(rows[:3])
            middle = 3 + (len(rows) - 3) // 2
            upper = set(rows[:middle])
            lower = headers | set(rows[middle:])
            groups = [
                [cell for cell in candidate.cells if cell.row in selected]
                for selected in (upper, lower)
            ]
        elif len(columns) > 1:
            middle_column = columns[len(columns) // 2]
            groups = [
                [cell for cell in candidate.cells if cell.column < middle_column],
                [cell for cell in candidate.cells if cell.column >= middle_column],
            ]
        elif len(rows) > 1:
            # An oversized header cannot be repeated indefinitely. Keep all
            # evidence in separate fragments rather than silently truncating it.
            middle_row = rows[len(rows) // 2]
            groups = [
                [cell for cell in candidate.cells if cell.row < middle_row],
                [cell for cell in candidate.cells if cell.row >= middle_row],
            ]
        else:
            raise ExcelTileBudgetError(
                f"Excel cell {candidate.sheet_name}!{candidate.cells[0].coordinate} "
                f"exceeds native token budget {budget}; cannot split a single cell."
            )
        result: list[ExcelRegion] = []
        for index, cells in enumerate(groups, 1):
            left, right = min(c.column for c in cells), max(c.column for c in cells)
            top, bottom = min(c.row for c in cells), max(c.row for c in cells)
            cell_range = f"{_column_letter(left)}{top}:{_column_letter(right)}{bottom}"
            child = candidate.model_copy(update={
                "region_id": f"{candidate.region_id}_b{index}",
                "parent_region_id": region.region_id,
                "is_tile": True,
                "min_row": top, "max_row": bottom,
                "min_column": left, "max_column": right,
                "cell_range": cell_range, "cells": cells,
                "native_text": _region_native_text(
                    sheet_name=candidate.sheet_name, cell_range=cell_range, cells=cells,
                ),
            })
            result.extend(refine(child))
        return result

    tiles = [
        child for tile in _split_region_by_dimensions(region, settings)
        for child in refine(tile)
    ]
    if len(tiles) > 1:
        logger.info(
            "[Excel] %s: %d tiles, peak native estimate=%d tokens (budget=%d).",
            region.region_id, len(tiles),
            max(_estimate_native_tokens(tile.native_text) for tile in tiles), budget,
        )
    return tiles


def split_excel_regions_into_tiles(
    survey: ExcelWorkbookSurvey, settings: Settings | None = None
) -> ExcelWorkbookSurvey:
    """Return a survey whose oversized regions are represented by extraction tiles."""
    resolved_settings = settings or get_settings()
    updated_sheets = []
    for sheet in survey.sheets:
        if not sheet.visible:
            updated_sheets.append(sheet)
            continue
        tiled_regions = [
            tile
            for region in sheet.regions
            for tile in _split_region_into_tiles(region, resolved_settings)
        ]
        updated_sheets.append(sheet.model_copy(update={"regions": tiled_regions}))
    return survey.model_copy(update={"sheets": updated_sheets})


def _combine_tile_regions(regions: list[ExcelRegion]) -> ExcelRegion:
    """Combine tile evidence back into one logical region by cell coordinate."""
    unique_cells = {
        item.coordinate: item
        for region in regions
        for item in region.cells
    }
    return regions[0].model_copy(
        update={
            "region_id": regions[0].parent_region_id or regions[0].region_id,
            "cell_range": (
                f"{_column_letter(min(item.column for item in unique_cells.values()))}"
                f"{min(item.row for item in unique_cells.values())}:"
                f"{_column_letter(max(item.column for item in unique_cells.values()))}"
                f"{max(item.row for item in unique_cells.values())}"
            ),
            "min_row": min(item.row for item in unique_cells.values()),
            "max_row": max(item.row for item in unique_cells.values()),
            "min_column": min(item.column for item in unique_cells.values()),
            "max_column": max(item.column for item in unique_cells.values()),
            "cells": list(unique_cells.values()),
            "native_text": "",
            "parent_region_id": None,
            "tile_row_index": 0,
            "tile_column_index": 0,
            "is_tile": False,
        }
    )


def _native_region_markdown(region: ExcelRegion) -> str:
    """Create one deterministic Markdown table from native cell evidence."""
    cells = {(item.row, item.column): item for item in region.cells}
    rows = sorted({row for row, _ in cells})
    columns = sorted({column for _, column in cells})
    if not rows or not columns:
        return ""
    header_rows = rows[: min(3, len(rows))]
    if len(header_rows) > 1 and sum(bool(cells.get((header_rows[0], col))) for col in columns) <= 1:
        header_rows = header_rows[1:]
    data_rows = [row for row in rows if row not in header_rows]

    def value(row: int, column: int) -> str:
        item = cells.get((row, column))
        if item is None:
            return ""
        raw = item.value
        if raw in (None, "") and item.formula:
            raw = f"={item.formula}"
        return str(raw or "").replace("|", "\\|").replace("\n", " ").strip()

    headers = []
    for column in columns:
        parts = [value(row, column) for row in header_rows]
        parts = [part for part in parts if part]
        headers.append(" ".join(dict.fromkeys(parts)) or _column_letter(column))
    lines = [
        "| " + " | ".join(headers) + " |",
        "| " + " | ".join("---" for _ in headers) + " |",
    ]
    for row in data_rows:
        lines.append("| " + " | ".join(value(row, column) for column in columns) + " |")
    return "\n".join(lines)


def _convert_workbook_to_xlsx_copy(
    source: Path,
    output_dir: Path,
) -> Path:
    """Buka, hitung ulang, dan simpan salinan XLSX sementara untuk survei."""
    soffice = _find_libreoffice_binary()
    if not soffice:
        raise RuntimeError("LibreOffice diperlukan untuk menyiapkan salinan survei XLSX.")
    output_dir.mkdir(parents=True, exist_ok=True)
    profile = output_dir / "lo_profile_xlsx"
    result = subprocess.run(
        [
            soffice,
            f"-env:UserInstallation={profile.as_uri()}",
            "--headless",
            "--convert-to",
            "xlsx",
            "--outdir",
            str(output_dir),
            str(source),
        ],
        capture_output=True,
        text=True,
        timeout=180,
        check=False,
    )
    target = output_dir / f"{source.stem}.xlsx"
    if result.returncode != 0 or not target.exists():
        raise RuntimeError(
            f"Konversi workbook ke XLSX gagal: {result.stderr.strip()}"
        )
    return target


def _convert_excel_regions_to_pdf(
    excel_path: Path,
    pdf_path: Path,
    profile_dir: Path,
    regions: list[ExcelRegion],
) -> bool:
    """Ekspor setiap region sebagai PDF terpisah lalu gabungkan sesuai urutan survei."""
    soffice = _find_libreoffice_binary()
    if not soffice or not regions:
        return False
    uno_python = _find_uno_python(soffice)
    if uno_python is None:
        return False

    profile_dir.mkdir(parents=True, exist_ok=True)
    region_dir = profile_dir.parent / "region_pdfs"
    region_dir.mkdir(parents=True, exist_ok=True)
    spec_path = profile_dir.parent / "excel_regions.json"
    spec_path.write_text(
        json.dumps(
            [
                {
                    "sheet_index": region.sheet_index,
                    "min_row": region.min_row - 1,
                    "max_row": region.max_row - 1,
                    "min_column": region.min_column - 1,
                    "max_column": region.max_column - 1,
                    "included_rows": (
                        sorted({cell.row - 1 for cell in region.cells})
                        if region.is_tile else None
                    ),
                    "output": str(region_dir / f"{index:04d}.pdf"),
                }
                for index, region in enumerate(regions, start=1)
            ],
            ensure_ascii=False,
        ),
        encoding="utf-8",
    )

    try:
        with socket.socket() as probe:
            probe.bind(("127.0.0.1", 0))
            port = probe.getsockname()[1]
    except OSError as exc:
        logger.warning("Tidak dapat membuka port lokal untuk render region Excel: %s", exc)
        return False

    listener = subprocess.Popen(
        [
            soffice,
            f"-env:UserInstallation={profile_dir.as_uri()}",
            "--headless",
            "--norestore",
            "--nodefault",
            "--nofirststartwizard",
            f"--accept=socket,host=127.0.0.1,port={port};urp;StarOffice.ComponentContext",
        ],
        stdout=subprocess.DEVNULL,
        stderr=subprocess.DEVNULL,
    )
    uno_script = r"""
import json
import sys
import time
import uno
from com.sun.star.beans import PropertyValue


def prop(name, value):
    item = PropertyValue()
    item.Name = name
    item.Value = value
    return item


port = int(sys.argv[1])
source_url = uno.systemPathToFileUrl(sys.argv[2])
specs = json.load(open(sys.argv[3], "r", encoding="utf-8"))
local_context = uno.getComponentContext()
resolver = local_context.ServiceManager.createInstanceWithContext(
    "com.sun.star.bridge.UnoUrlResolver", local_context
)
context = None
for _ in range(60):
    try:
        context = resolver.resolve(
            f"uno:socket,host=127.0.0.1,port={port};urp;StarOffice.ComponentContext"
        )
        break
    except Exception:
        time.sleep(0.25)
if context is None:
    raise RuntimeError("Tidak dapat terhubung ke listener LibreOffice UNO")

desktop = context.ServiceManager.createInstanceWithContext(
    "com.sun.star.frame.Desktop", context
)
document = desktop.loadComponentFromURL(
    source_url, "_blank", 0, (prop("Hidden", True),)
)
if document is None:
    raise RuntimeError("LibreOffice tidak dapat membuka workbook")

try:
    sheets = document.getSheets()
    page_styles = document.getStyleFamilies().getByName("PageStyles")
    changed_rows = []
    for spec in specs:
        for row, was_visible in changed_rows:
            row.setPropertyValue("IsVisible", was_visible)
        changed_rows = []
        for sheet_index in range(sheets.getCount()):
            sheet = sheets.getByIndex(sheet_index)
            sheet.setPrintAreas(())
            try:
                sheet.setPropertyValue("AutomaticPrintArea", False)
            except Exception:
                pass
        selected = sheets.getByIndex(int(spec["sheet_index"]))
        included_rows = spec.get("included_rows")
        if included_rows is not None:
            included_rows = set(included_rows)
            sheet_rows = selected.getRows()
            for row_index in range(int(spec["min_row"]), int(spec["max_row"]) + 1):
                if row_index not in included_rows:
                    row = sheet_rows.getByIndex(row_index)
                    changed_rows.append((row, row.getPropertyValue("IsVisible")))
                    row.setPropertyValue("IsVisible", False)
        address = uno.createUnoStruct("com.sun.star.table.CellRangeAddress")
        address.Sheet = int(spec["sheet_index"])
        address.StartColumn = int(spec["min_column"])
        address.EndColumn = int(spec["max_column"])
        address.StartRow = int(spec["min_row"])
        address.EndRow = int(spec["max_row"])
        selected.setPrintAreas((address,))
        page_style = page_styles.getByName(selected.getPropertyValue("PageStyle"))
        page_style.setPropertyValue("ScaleToPagesX", 1)
        page_style.setPropertyValue("ScaleToPagesY", 1)
        document.calculateAll()
        document.storeToURL(
            uno.systemPathToFileUrl(spec["output"]),
            (prop("FilterName", "calc_pdf_Export"), prop("Overwrite", True)),
        )
finally:
    document.close(True)
    desktop.terminate()
"""
    try:
        subprocess.run(
            [uno_python, "-c", uno_script, str(port), str(excel_path), str(spec_path)],
            capture_output=True,
            text=True,
            timeout=max(180, 60 * len(regions)),
            check=True,
        )
        import pymupdf

        merged = pymupdf.open()
        try:
            for index in range(1, len(regions) + 1):
                region_pdf = region_dir / f"{index:04d}.pdf"
                if not region_pdf.exists():
                    raise RuntimeError(f"PDF region tidak ditemukan: {region_pdf}")
                source_pdf = pymupdf.open(region_pdf)
                try:
                    if len(source_pdf) != 1:
                        raise RuntimeError(
                            f"Region {index} menghasilkan {len(source_pdf)} halaman."
                        )
                    merged.insert_pdf(source_pdf)
                finally:
                    source_pdf.close()
            merged.save(pdf_path)
        finally:
            merged.close()
        return pdf_path.exists() and pdf_page_count(pdf_path) == len(regions)
    except (OSError, subprocess.CalledProcessError, subprocess.TimeoutExpired, RuntimeError) as exc:
        logger.warning("Render region Excel gagal; gunakan render seluruh sheet: %s", exc)
        return False
    finally:
        listener.terminate()
        try:
            listener.wait(timeout=10)
        except subprocess.TimeoutExpired:
            listener.kill()
            listener.wait()
        time.sleep(0.5)


def plan_excel_sheet_previews(survey: ExcelWorkbookSurvey) -> list[ExcelRegion]:
    """Rencanakan gambar pembanding tanpa menambah pekerjaan OCR/VLM."""
    from openpyxl.utils.cell import range_boundaries

    sheets = {sheet.name: sheet for sheet in survey.sheets}
    previews: list[ExcelRegion] = []
    for sheet_name in survey.extraction_order:
        sheet = sheets[sheet_name]
        if not sheet.visible:
            continue
        min_col, min_row, max_col, max_row = (
            range_boundaries(sheet.used_range) if sheet.used_range else (1, 1, 1, 1)
        )
        rows = max_row - min_row + 1
        columns = max_col - min_col + 1
        # Sheet data besar dipaginasi supaya hasil cetaknya tetap terbaca.
        row_step = 50 if sheet.role == "detail" or rows > 80 or columns > 30 else rows
        col_step = 20 if sheet.role == "detail" or rows > 80 or columns > 30 else columns
        sheet_page = 0
        for row_start in range(min_row, max_row + 1, row_step):
            row_end = min(row_start + row_step - 1, max_row)
            for col_start in range(min_col, max_col + 1, col_step):
                col_end = min(col_start + col_step - 1, max_col)
                sheet_page += 1
                previews.append(
                    ExcelRegion(
                        region_id=f"sheet_preview_{sheet.index + 1:03d}_{sheet_page:03d}",
                        sheet_name=sheet.name,
                        sheet_index=sheet.index,
                        cell_range=(
                            f"{_column_letter(col_start)}{row_start}:"
                            f"{_column_letter(col_end)}{row_end}"
                        ),
                        min_row=row_start,
                        max_row=row_end,
                        min_column=col_start,
                        max_column=col_end,
                        title=sheet.name,
                        kind="mixed",
                        render_strategy="visual",
                        persist_native=False,
                    )
                )
    return previews


def render_excel_sheet_previews(
    excel_path: Path,
    survey: ExcelWorkbookSurvey,
    output_root: Path,
    temporary_root: Path,
) -> list[dict[str, Any]]:
    """Simpan gambar tiap sheet dan pemetaan gambarnya ke section Markdown."""
    regions = plan_excel_sheet_previews(survey)
    if not regions:
        return []
    pdf_path = temporary_root / "sheet_previews.pdf"
    if not _convert_excel_regions_to_pdf(
        excel_path, pdf_path, temporary_root / "lo_sheet_preview_profile", regions
    ):
        logger.warning("[Excel] Gambar pembanding per sheet tidak dapat dirender.")
        return []

    output_root.mkdir(parents=True, exist_ok=True)
    staged_dir = Path(tempfile.mkdtemp(prefix=".sheet_previews_", dir=output_root))
    try:
        images = pdf_to_images(pdf_path, staged_dir, dpi=160)
        if len(images) != len(regions):
            raise RuntimeError(
                f"Jumlah gambar sheet {len(images)} tidak sesuai rencana {len(regions)}."
            )
        from PIL import Image

        for image_path in images:
            with Image.open(image_path) as original:
                grayscale = original.convert("L")
                bounds = grayscale.point(lambda value: 255 if value < 245 else 0).getbbox()
                if bounds:
                    margin = 24
                    crop = (
                        max(0, bounds[0] - margin),
                        max(0, bounds[1] - margin),
                        min(original.width, bounds[2] + margin),
                        min(original.height, bounds[3] + margin),
                    )
                    original.crop(crop).save(image_path)
        page_counts: dict[str, int] = {}
        entries: list[dict[str, Any]] = []
        for image_path, region in zip(images, regions, strict=True):
            page_counts[region.sheet_name] = page_counts.get(region.sheet_name, 0) + 1
            entries.append(
                {
                    "sheet_name": region.sheet_name,
                    "sheet_index": region.sheet_index,
                    "sheet_page": page_counts[region.sheet_name],
                    "cell_range": region.cell_range,
                    "image": image_path.name,
                }
            )
        (staged_dir / "manifest.json").write_text(
            json.dumps({"source_file": str(excel_path), "pages": entries}, ensure_ascii=False, indent=2),
            encoding="utf-8",
        )
        final_dir = output_root / "sheet_previews"
        if final_dir.exists():
            shutil.rmtree(final_dir)
        staged_dir.replace(final_dir)
        logger.info("[Excel] %d gambar pembanding untuk %d sheet tersimpan.", len(entries), len(page_counts))
        return entries
    finally:
        if staged_dir.exists():
            shutil.rmtree(staged_dir)


def _header_candidate_score(
    nonempty: list[ExcelCellEvidence],
    following: list[ExcelCellEvidence],
    *,
    width: int,
    text_count: int,
) -> float:
    following_numeric = sum(
        isinstance(cell.value, (int, float)) and not isinstance(cell.value, bool)
        for cell in following
    )
    coverage = len(nonempty) / width
    normalized_headers = {str(cell.value).strip().casefold() for cell in nonempty}
    uniqueness = len(normalized_headers) / len(nonempty)
    immediate_body_bonus = 3.0 if len(following) >= 2 else 0.0
    formula_penalty = sum(bool(cell.formula) for cell in nonempty) * 0.75
    long_text_penalty = sum(
        len(str(cell.value).strip()) > 80 for cell in nonempty
    ) * 0.5
    return (
        text_count
        + 2.0 * coverage
        + uniqueness
        + immediate_body_bonus
        + min(2.0, following_numeric * 0.5)
        - formula_penalty
        - long_text_penalty
    )


def _detect_header_row(region: ExcelRegion) -> tuple[int, list[str]] | None:
    by_row: dict[int, list[ExcelCellEvidence]] = {}
    for cell in region.cells:
        by_row.setdefault(cell.row, []).append(cell)
    candidates: list[tuple[float, int, list[str]]] = []
    width = max(1, region.max_column - region.min_column + 1)
    candidate_rows = sorted(by_row)[:12]
    for row_number in candidate_rows:
        cells = sorted(by_row[row_number], key=lambda item: item.column)
        text_count = sum(
            isinstance(cell.value, str) and bool(cell.value.strip()) for cell in cells
        )
        nonempty = [cell for cell in cells if cell.value not in (None, "")]
        if text_count >= 2 and len(nonempty) >= 2:
            headers_by_column = {cell.column: str(cell.value).strip() for cell in nonempty}
            headers = [
                headers_by_column.get(column, f"column_{column - region.min_column + 1}")
                for column in range(region.min_column, region.max_column + 1)
            ]
            following = [
                cell
                for cell in by_row.get(row_number + 1, [])
                if cell.value not in (None, "")
            ]
            score = _header_candidate_score(
                nonempty,
                following,
                width=width,
                text_count=text_count,
            )
            candidates.append((score, row_number, headers))
    if not candidates:
        return None
    _, header_row, headers = max(candidates, key=lambda item: (item[0], item[1]))
    return header_row, headers


_EXCEL_METADATA_COLUMNS = frozenset(
    {"source_file", "sheet_name", "region_id", "period", "source_row", "row_role"}
)


def _unique_sql_headers(headers: list[str]) -> list[str]:
    seen: dict[str, int] = {}
    result: list[str] = []
    for index, header in enumerate(headers, start=1):
        base = _slug(header, fallback=f"column_{index}")
        if base in _EXCEL_METADATA_COLUMNS:
            # Hindari bentrok dengan kolom metadata tabel native.
            base = f"{base}_value"
        seen[base] = seen.get(base, 0) + 1
        result.append(base if seen[base] == 1 else f"{base}_{seen[base]}")
    return result


def _infer_excel_sqlite_type(values: list[Any]) -> str:
    """Pilih afinitas SQLite tanpa mengubah teks/kode Excel menjadi angka."""
    non_empty = [value for value in values if value not in (None, "")]
    if not non_empty:
        return "TEXT"
    if not all(
        isinstance(value, (int, float)) and not isinstance(value, bool)
        for value in non_empty
    ):
        return "TEXT"
    if all(
        isinstance(value, int)
        or (isinstance(value, float) and value.is_integer())
        for value in non_empty
    ):
        return "INTEGER"
    return "REAL"


def _coerce_excel_sqlite_value(value: Any, sql_type: str) -> Any:
    """Pertahankan representasi tekstual stabil untuk kolom nonnumerik."""
    if value is None or sql_type != "TEXT":
        return value
    if isinstance(value, (datetime, date, datetime_time)):
        return value.isoformat()
    return str(value)


def _markdown_cell(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return str(value).replace("|", "\\|").replace("\n", "<br>").strip()


def _parameter_markdown(region: ExcelRegion) -> str:
    by_row: dict[int, dict[int, Any]] = {}
    for cell in region.cells:
        if cell.value not in (None, ""):
            by_row.setdefault(cell.row, {})[cell.column] = cell.value
    lines = []
    for row in sorted(by_row):
        label = _markdown_cell(by_row[row].get(region.min_column)).rstrip(":")
        value = _markdown_cell(by_row[row].get(region.max_column))
        lines.append(f"- **{label}**: {value}")
    note = (
        "Parameter/filter laporan."
        if region.referenced_by_formula
        else "Parameter/filter tampilan; tidak dirujuk formula mana pun."
    )
    return "\n".join([f"> {note}", *lines])


def _aggregate_note(region: ExcelRegion, rows: dict[int, str]) -> str:
    labels: list[str] = []
    for row in sorted(rows):
        label = next(
            (
                _markdown_cell(cell.value)
                for cell in sorted(region.cells, key=lambda item: item.column)
                if cell.row == row and isinstance(cell.value, str) and cell.value.strip()
            ),
            f"baris {row}",
        )
        labels.append(label)
    shown = ", ".join(labels[:8]) + (" ..." if len(labels) > 8 else "")
    return (
        f"> Baris agregat ({len(rows)}): {shown}. "
        "Jangan dijumlahkan bersama baris data (kolom SQLite `row_role`)."
    )


def _region_markdown(region: ExcelRegion, *, max_rows: int = 200) -> str:
    if region.semantic_role == "parameter":
        return _parameter_markdown(region)
    if region.kind == "text":
        values = [
            _markdown_cell(cell.value)
            for cell in sorted(region.cells, key=lambda item: (item.row, item.column))
            if cell.value not in (None, "")
        ]
        return "\n\n".join(dict.fromkeys(values))

    header = _detect_header_row(region)
    if header is None:
        return _region_native_text(
            sheet_name=region.sheet_name,
            cell_range=region.cell_range,
            cells=region.cells,
        )
    header_row, raw_headers = header
    headers = [header or f"Kolom {index}" for index, header in enumerate(raw_headers, 1)]
    values_by_coordinate = {
        (cell.row, cell.column): cell.value for cell in region.cells
    }
    rows: list[list[str]] = []
    for row_number in range(header_row + 1, region.max_row + 1):
        values = [
            _markdown_cell(values_by_coordinate.get((row_number, column)))
            for column in range(region.min_column, region.max_column + 1)
        ]
        if any(values):
            rows.append(values)
    visible_rows = rows[:max_rows]
    lines = [
        "| " + " | ".join(_markdown_cell(value) for value in headers) + " |",
        "| " + " | ".join("---" for _ in headers) + " |",
    ]
    lines.extend("| " + " | ".join(row) + " |" for row in visible_rows)
    aggregate_rows = _aggregate_row_roles(region, header_row)
    if aggregate_rows:
        lines.append("\n" + _aggregate_note(region, aggregate_rows))
    if len(rows) > max_rows:
        lines.append(
            f"\n> Tabel diringkas: {max_rows} dari {len(rows)} baris ditampilkan. "
            "Data lengkap tersedia pada artefak CSV/SQLite."
        )
    return "\n".join(lines)


def _chart_points(
    series: ExcelChartSeries,
) -> list[tuple[Any, float | int]]:
    return [
        (category, value)
        for category, value in zip(series.categories, series.values, strict=False)
        if category not in (None, "") and value is not None
    ]


def _chart_narrative(chart: ExcelChartEvidence) -> list[str]:
    narratives: list[str] = []
    for series in chart.series:
        points = _chart_points(series)
        if not points:
            continue
        date_points = [
            (category, value)
            for category, value in points
            if _category_family(category) == "date"
        ]
        if len(date_points) >= 2 and len(date_points) >= len(points) / 2:
            start_category, start_value = date_points[0]
            end_category, end_value = date_points[-1]
            max_category, max_value = max(date_points, key=lambda item: item[1])
            min_category, min_value = min(date_points, key=lambda item: item[1])
            change = end_value - start_value
            narratives.append(
                f"Seri **{series.name}** berubah dari {_markdown_cell(start_value)} "
                f"pada {_markdown_cell(start_category)} menjadi {_markdown_cell(end_value)} "
                f"pada {_markdown_cell(end_category)} (perubahan {_markdown_cell(change)}). "
                f"Nilai tertinggi {_markdown_cell(max_value)} pada "
                f"{_markdown_cell(max_category)} dan terendah {_markdown_cell(min_value)} "
                f"pada {_markdown_cell(min_category)}."
            )
        else:
            max_category, max_value = max(points, key=lambda item: item[1])
            min_category, min_value = min(points, key=lambda item: item[1])
            narratives.append(
                f"Pada seri **{series.name}**, nilai terbesar adalah "
                f"{_markdown_cell(max_value)} untuk {_markdown_cell(max_category)}, "
                f"sedangkan nilai terkecil adalah {_markdown_cell(min_value)} untuk "
                f"{_markdown_cell(min_category)}."
            )
    return narratives


def _chart_markdown(chart: ExcelChartEvidence, *, max_points: int = 200) -> str:
    max_length = max((len(series.categories) for series in chart.series), default=0)
    headers = ["Kategori", *(series.name for series in chart.series)]
    rows: list[list[str]] = []
    for index in range(min(max_length, max_points)):
        category = next(
            (
                series.categories[index]
                for series in chart.series
                if index < len(series.categories)
                and series.categories[index] not in (None, "")
            ),
            "",
        )
        values = [
            series.values[index] if index < len(series.values) else None
            for series in chart.series
        ]
        if category not in (None, "") or any(value is not None for value in values):
            rows.append([_markdown_cell(category), *map(_markdown_cell, values)])
    parts = [
        f"### Grafik: {chart.title}",
        "| " + " | ".join(_markdown_cell(value) for value in headers) + " |",
        "| " + " | ".join("---" for _ in headers) + " |",
        *("| " + " | ".join(row) + " |" for row in rows),
    ]
    narratives = _chart_narrative(chart)
    if narratives:
        parts.append("\n".join(f"> {item}" for item in narratives))
    if max_length > max_points:
        parts.append(
            f"> Grafik memiliki {max_length} titik; hanya {max_points} titik pertama "
            "ditampilkan di Markdown."
        )
    if chart.warnings:
        parts.append("\n".join(f"> Peringatan: {item}" for item in chart.warnings))
    return "\n".join(parts)


def _reference_cell_range(reference: str | None) -> str | None:
    if not reference or "!" not in reference:
        return None
    return reference.rsplit("!", 1)[1].replace("$", "")


def _reference_sheet_name(reference: str | None) -> str | None:
    if not reference or "!" not in reference:
        return None
    return reference.rsplit("!", 1)[0].strip("'").replace("''", "'")


def _range_contains(outer: str, inner: str) -> bool:
    try:
        from openpyxl.utils.cell import range_boundaries

        outer_min_col, outer_min_row, outer_max_col, outer_max_row = range_boundaries(
            outer
        )
        inner_min_col, inner_min_row, inner_max_col, inner_max_row = range_boundaries(
            inner
        )
    except (ImportError, TypeError, ValueError):
        return False
    return (
        outer_min_col <= inner_min_col <= inner_max_col <= outer_max_col
        and outer_min_row <= inner_min_row <= inner_max_row <= outer_max_row
    )


def _region_is_chart_source(region: ExcelRegion, source_ranges: set[str]) -> bool:
    """Region dianggap sumber grafik bila memuatnya atau mayoritas selnya dirujuk."""
    if any(_range_contains(region.cell_range, source) for source in source_ranges):
        return True
    source_bounds = [
        bounds
        for bounds in (_range_bounds(source) for source in source_ranges)
        if bounds is not None
    ]
    populated = [
        (cell.row, cell.column) for cell in region.cells if cell.value not in (None, "")
    ]
    if not populated or not source_bounds:
        return False
    covered = sum(
        any(_coordinate_in_bounds(coordinate, bounds) for bounds in source_bounds)
        for coordinate in populated
    )
    return covered / len(populated) >= 0.5


def _combine_markdown_table_tiles(regions: list[ExcelRegion]) -> list[ExcelRegion]:
    """Render each tiled native table as one logical Markdown table."""
    tile_groups: dict[str, list[ExcelRegion]] = {}
    for region in regions:
        if region.kind == "table" and region.parent_region_id:
            tile_groups.setdefault(region.parent_region_id, []).append(region)
    if not tile_groups:
        return regions

    combined: list[ExcelRegion] = []
    emitted_parents: set[str] = set()
    for region in regions:
        parent_id = region.parent_region_id
        if region.kind == "table" and parent_id in tile_groups:
            if parent_id in emitted_parents:
                continue
            combined.append(_combine_tile_regions(tile_groups[parent_id]))
            emitted_parents.add(parent_id)
        else:
            combined.append(region)
    return combined


def _sheet_markdown(
    sheet: ExcelSheetSurvey,
    *,
    artifacts: list[ExcelNativeArtifact],
    visual_markdown: str | None = None,
) -> str:
    parts = [f"## Sheet: {sheet.name}"]
    if sheet.role == "detail":
        sheet_artifacts = [
            artifact for artifact in artifacts if artifact.sheet_name == sheet.name
        ]
        if sheet_artifacts:
            data_rows = sum(artifact.row_count for artifact in sheet_artifacts)
            columns = next(
                (artifact.columns for artifact in sheet_artifacts if artifact.columns),
                [],
            )
        else:
            largest = max(
                (region for region in sheet.regions if region.kind == "table"),
                key=lambda region: len(region.cells),
                default=None,
            )
            header = _detect_header_row(largest) if largest else None
            data_rows = (
                max(0, largest.max_row - header[0])
                if largest is not None and header is not None
                else sheet.row_count
            )
            columns = header[1] if header else []
        parts.append(f"Data detail terdiri dari {data_rows} baris dan {len(columns)} kolom.")
        if columns:
            parts.append("Kolom: " + ", ".join(f"`{column}`" for column in columns))
        for artifact in sheet_artifacts:
            location = artifact.csv_path or artifact.table_name
            parts.append(
                f"- Data lengkap: `{location}` ({artifact.row_count} baris)."
            )
        return "\n\n".join(parts)
    if sheet.role == "support":
        dependencies = ", ".join(sheet.dependencies) or "tidak ada"
        parts.append(
            "Sheet pendukung dipertahankan dalam manifest untuk audit formula. "
            f"Sumber yang dirujuk: {dependencies}."
        )
        for region in _combine_markdown_table_tiles(list(sheet.regions)):
            if region.kind != "table" or not region.persist_native:
                continue
            rendered = _region_markdown(region)
            if rendered:
                heading = region.title or region.cell_range
                parts.append(f"### {heading}\n\n{rendered}")
        return "\n\n".join(parts)

    chart_source_ranges = {
        cell_range
        for chart in sheet.charts
        for series in chart.series
        for reference, cell_range in (
            (
                series.category_reference,
                _reference_cell_range(series.category_reference),
            ),
            (series.value_reference, _reference_cell_range(series.value_reference)),
        )
        if cell_range and _reference_sheet_name(reference) == sheet.name
    }
    for region in _combine_markdown_table_tiles(list(sheet.regions)):
        if region.kind == "mixed":
            continue
        if sheet.role == "dashboard" and _region_is_chart_source(
            region, chart_source_ranges
        ):
            continue
        rendered = _region_markdown(region)
        if rendered:
            if region.kind == "table" and region.title:
                parts.append(f"### {region.title}\n\n{rendered}")
            else:
                parts.append(rendered)
    parts.extend(_chart_markdown(chart) for chart in sheet.charts)
    if visual_markdown and visual_markdown.strip():
        parts.append(f"### Konteks visual\n\n{visual_markdown.strip()}")
    return "\n\n".join(parts)


def _workbook_title(survey: ExcelWorkbookSurvey, fallback: str) -> str:
    ordered_names = survey.extraction_order or [sheet.name for sheet in survey.sheets]
    sheets = {sheet.name: sheet for sheet in survey.sheets}
    for name in ordered_names:
        sheet = sheets[name]
        if sheet.role != "dashboard":
            continue
        for region in sheet.regions:
            for cell in sorted(region.cells, key=lambda item: (item.row, item.column)):
                if isinstance(cell.value, str) and cell.value.strip():
                    return cell.value.strip()
    return fallback


def compose_excel_markdown(
    survey: ExcelWorkbookSurvey,
    *,
    fallback_title: str,
    artifacts: list[ExcelNativeArtifact] | None = None,
    visual_by_sheet: dict[str, str] | None = None,
) -> tuple[str, list[tuple[str, str]]]:
    artifacts = artifacts or []
    visual_by_sheet = visual_by_sheet or {}
    title = _workbook_title(survey, fallback_title)
    sheets_by_name = {sheet.name: sheet for sheet in survey.sheets}
    sections: list[tuple[str, str]] = []
    extraction_order = survey.extraction_order or [
        sheet.name for sheet in survey.sheets if sheet.visible
    ]
    for sheet_name in extraction_order:
        sheet = sheets_by_name[sheet_name]
        markdown = _sheet_markdown(
            sheet,
            artifacts=artifacts,
            visual_markdown=visual_by_sheet.get(sheet_name),
        )
        sections.append((sheet_name, markdown))
    if survey.warnings:
        warning_text = "## Peringatan Workbook\n\n" + "\n".join(
            f"- {warning}" for warning in survey.warnings
        )
        sections.append(("Peringatan Workbook", warning_text))
    full_markdown = f"# {title}\n\n" + "\n\n---\n\n".join(
        markdown for _, markdown in sections if markdown.strip()
    )
    return full_markdown.strip(), sections


def plan_excel_visual_regions(
    survey: ExcelWorkbookSurvey,
    *,
    settings: Settings | None = None,
) -> list[ExcelRegion]:
    """Pilih unit visual yang benar-benar memerlukan pembacaan VLM."""
    resolved_settings = settings or get_settings()
    planned: list[ExcelRegion] = []
    for sheet in survey.sheets:
        if not sheet.visible:
            continue
        if sheet.role == "dashboard" and sheet.used_range:
            cells_by_coordinate = {
                cell.coordinate: cell
                for region in sheet.regions
                for cell in region.cells
            }
            min_row = min((region.min_row for region in sheet.regions), default=1)
            max_row = max((region.max_row for region in sheet.regions), default=1)
            min_column = min(
                (region.min_column for region in sheet.regions), default=1
            )
            max_column = max(
                (region.max_column for region in sheet.regions), default=1
            )
            planned.append(
                ExcelRegion(
                    region_id=f"sheet_{sheet.index + 1:03d}_dashboard",
                    sheet_name=sheet.name,
                    sheet_index=sheet.index,
                    cell_range=sheet.used_range,
                    min_row=min_row,
                    max_row=max_row,
                    min_column=min_column,
                    max_column=max_column,
                    title=sheet.name,
                    kind="mixed",
                    render_dpi=resolved_settings.excel_base_dpi,
                    requires_vlm_reading=True,
                    render_strategy="visual",
                    persist_native=False,
                    native_text=_sheet_markdown(sheet, artifacts=[]),
                    cells=list(cells_by_coordinate.values()),
                )
            )
            continue
        if sheet.role in {"detail", "support"}:
            continue
        planned.extend(
            region
            for region in sheet.regions
            if region.render_strategy in {"visual", "hybrid"}
            and (region.requires_vlm_reading or region.kind == "mixed")
        )
    return planned


def _describe_excel_native_artifacts(
    survey: ExcelWorkbookSurvey,
    db_path: Path,
    table_names: list[str],
    csv_paths: list[Path],
) -> list[ExcelNativeArtifact]:
    csv_by_stem = {path.stem: path for path in csv_paths}
    artifacts: list[ExcelNativeArtifact] = []
    with closing(sqlite3.connect(db_path)) as connection:
        for table_name in table_names:
            quoted = table_name.replace('"', '""')
            columns = [
                row[1]
                for row in connection.execute(
                    f'PRAGMA table_info("{quoted}")'
                ).fetchall()
                if row[1]
                not in {
                    "source_file",
                    "sheet_name",
                    "region_id",
                    "period",
                    "source_row",
                    "row_role",
                }
            ]
            has_row_role = any(
                row[1] == "row_role"
                for row in connection.execute(f'PRAGMA table_info("{quoted}")')
            )
            row_count = connection.execute(
                f'SELECT COUNT(*) FROM "{quoted}" WHERE source_file = ?',
                (survey.source_file,),
            ).fetchone()[0]
            sheet_row = connection.execute(
                f'SELECT sheet_name FROM "{quoted}" WHERE source_file = ? LIMIT 1',
                (survey.source_file,),
            ).fetchone()
            aggregate_rows = (
                connection.execute(
                    f'SELECT COUNT(*) FROM "{quoted}" '
                    "WHERE source_file = ? AND row_role != 'data'",
                    (survey.source_file,),
                ).fetchone()[0]
                if has_row_role
                else 0
            )
            artifacts.append(
                ExcelNativeArtifact(
                    sheet_name=sheet_row[0] if sheet_row else "",
                    table_name=table_name,
                    row_count=row_count,
                    aggregate_row_count=aggregate_rows,
                    columns=columns,
                    csv_path=(
                        str(csv_by_stem[table_name])
                        if table_name in csv_by_stem
                        else None
                    ),
                )
            )
    return artifacts


def persist_excel_native_data(
    survey: ExcelWorkbookSurvey,
    db_path: str | Path,
) -> list[str]:
    """Simpan bukti sel dan dataset berulang tanpa mengandalkan hasil OCR/VLM."""
    target = Path(db_path).resolve()
    target.parent.mkdir(parents=True, exist_ok=True)
    created_tables: list[str] = []
    logical_sheets: dict[str, list[ExcelRegion]] = {}
    for sheet in survey.sheets:
        groups: dict[str, list[ExcelRegion]] = {}
        for region in sheet.regions:
            groups.setdefault(region.parent_region_id or region.region_id, []).append(region)
        logical_sheets[sheet.name] = [
            _combine_tile_regions(group) if len(group) > 1 and group[0].is_tile else group[0]
            for group in groups.values()
        ]
    grouped_regions: dict[tuple[str, tuple[str, ...]], list[tuple[ExcelRegion, int]]] = {}

    with closing(sqlite3.connect(target)) as connection:
        connection.execute(
            """
            CREATE TABLE IF NOT EXISTS excel_regions (
                source_file TEXT NOT NULL,
                sheet_name TEXT NOT NULL,
                region_id TEXT NOT NULL,
                cell_range TEXT NOT NULL,
                title TEXT,
                period TEXT,
                kind TEXT NOT NULL,
                render_dpi INTEGER NOT NULL,
                requires_vlm_reading INTEGER NOT NULL,
                PRIMARY KEY (source_file, region_id)
            )
            """
        )
        connection.execute(
            """
            CREATE TABLE IF NOT EXISTS excel_cells (
                source_file TEXT NOT NULL,
                sheet_name TEXT NOT NULL,
                region_id TEXT NOT NULL,
                coordinate TEXT NOT NULL,
                row_number INTEGER NOT NULL,
                column_number INTEGER NOT NULL,
                value_text TEXT,
                value_number REAL,
                formula TEXT,
                number_format TEXT,
                data_type TEXT,
                merged_range TEXT,
                PRIMARY KEY (source_file, sheet_name, region_id, coordinate)
            )
            """
        )
        connection.execute("DELETE FROM excel_regions WHERE source_file = ?", (survey.source_file,))
        connection.execute("DELETE FROM excel_cells WHERE source_file = ?", (survey.source_file,))

        for sheet in survey.sheets:
            for region in logical_sheets[sheet.name]:
                connection.execute(
                    """
                    INSERT INTO excel_regions VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?)
                    """,
                    (
                        survey.source_file,
                        region.sheet_name,
                        region.region_id,
                        region.cell_range,
                        region.title,
                        region.period,
                        region.kind,
                        region.render_dpi,
                        int(region.requires_vlm_reading),
                    ),
                )
                for cell in region.cells:
                    numeric = (
                        float(cell.value)
                        if isinstance(cell.value, (int, float))
                        and not isinstance(cell.value, bool)
                        else None
                    )
                    connection.execute(
                        """
                        INSERT INTO excel_cells VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
                        """,
                        (
                            survey.source_file,
                            region.sheet_name,
                            region.region_id,
                            cell.coordinate,
                            cell.row,
                            cell.column,
                            None if cell.value is None else str(cell.value),
                            numeric,
                            cell.formula,
                            cell.number_format,
                            cell.data_type,
                            cell.merged_range,
                        ),
                    )

                header = _detect_header_row(region) if region.persist_native else None
                if header:
                    header_row, raw_headers = header
                    sql_headers = _unique_sql_headers(raw_headers)
                    grouped_regions.setdefault(
                        (region.sheet_name, tuple(sql_headers)), []
                    ).append((region, header_row))

        for (sheet_name, headers), group in grouped_regions.items():
            signature = hashlib.sha1(
                (sheet_name + "|" + "|".join(headers)).encode("utf-8")
            ).hexdigest()[:8]
            table_name = f"excel_{_slug(sheet_name)}_{signature}"
            prepared_rows: list[tuple[ExcelRegion, int, str, list[Any]]] = []
            for region, header_row in group:
                values_by_coordinate = {
                    (cell.row, cell.column): cell.value for cell in region.cells
                }
                row_roles = _aggregate_row_roles(region, header_row)
                for row in range(header_row + 1, region.max_row + 1):
                    values = [
                        values_by_coordinate.get((row, column))
                        for column in range(region.min_column, region.max_column + 1)
                    ]
                    if any(value not in (None, "") for value in values):
                        prepared_rows.append(
                            (region, row, row_roles.get(row, "data"), values)
                        )

            column_types = [
                _infer_excel_sqlite_type(
                    [values[index] for *_, values in prepared_rows]
                )
                for index in range(len(headers))
            ]
            quoted_columns = ", ".join(
                f'"{header}" {sql_type}'
                for header, sql_type in zip(headers, column_types, strict=True)
            )
            connection.execute(
                f"""
                CREATE TABLE IF NOT EXISTS "{table_name}" (
                    source_file TEXT NOT NULL,
                    sheet_name TEXT NOT NULL,
                    region_id TEXT NOT NULL,
                    period TEXT,
                    source_row INTEGER NOT NULL,
                    row_role TEXT NOT NULL DEFAULT 'data',
                    {quoted_columns}
                )
                """
            )
            existing_columns = {
                row[1]
                for row in connection.execute(f'PRAGMA table_info("{table_name}")')
            }
            if "row_role" not in existing_columns:
                # Database lama dibuat sebelum kolom row_role tersedia.
                connection.execute(
                    f'ALTER TABLE "{table_name}" '
                    "ADD COLUMN row_role TEXT NOT NULL DEFAULT 'data'"
                )
            connection.execute(
                f'DELETE FROM "{table_name}" WHERE source_file = ?',
                (survey.source_file,),
            )
            insert_columns = [
                "source_file",
                "sheet_name",
                "region_id",
                "period",
                "source_row",
                "row_role",
                *headers,
            ]
            column_sql = ", ".join(f'"{column}"' for column in insert_columns)
            placeholders = ", ".join("?" for _ in insert_columns)
            for region, row, row_role, values in prepared_rows:
                stored_values = [
                    _coerce_excel_sqlite_value(value, sql_type)
                    for value, sql_type in zip(values, column_types, strict=True)
                ]
                connection.execute(
                    f'INSERT INTO "{table_name}" ({column_sql}) VALUES ({placeholders})',
                    (
                        survey.source_file,
                        region.sheet_name,
                        region.region_id,
                        region.period,
                        row,
                        row_role,
                        *stored_values,
                    ),
                )
            created_tables.append(table_name)
        connection.commit()
    return created_tables


def convert_excel_for_extraction(
    excel_path: str | Path,
    output_dir: str | Path,
) -> Path:
    """Render rencana visual adaptif, dengan konversi penuh sebagai fallback."""
    source = Path(excel_path).resolve()
    target_dir = Path(output_dir).resolve()
    target_dir.mkdir(parents=True, exist_ok=True)
    settings = get_settings()
    target = target_dir / f"{source.stem}.pdf"
    try:
        survey = survey_excel_workbook(source, settings=settings)
        visual_regions = plan_excel_visual_regions(survey, settings=settings)
        if visual_regions and _convert_excel_regions_to_pdf(
            source,
            target,
            target_dir / "lo_region_profile",
            visual_regions,
        ):
            return target
    except (OSError, ValueError, RuntimeError, TypeError) as exc:
        logger.warning("Survei adaptif gagal; gunakan konversi penuh: %s", exc)
    return convert_excel_to_pdf(source, target_dir)


def excel_page_count(excel_path: str | Path) -> int:
    """Hitung jumlah unit visual dari rencana ekstraksi adaptif."""
    with tempfile.TemporaryDirectory(prefix="excel_count_") as tmp:
        pdf_path = convert_excel_for_extraction(excel_path, tmp)
        return pdf_page_count(pdf_path)


def _visual_regions_need_vlm_tables(
    survey: ExcelWorkbookSurvey | None,
    visual_regions: list[ExcelRegion],
) -> bool:
    """Tentukan apakah tabel hasil VLM perlu di-ingest ke SQLite.

    Untuk Excel, angka tabel dan grafik sudah tersedia dari struktur workbook dan
    disimpan oleh ``persist_excel_native_data``. Tabel hasil pembacaan visual
    hanya dibutuhkan bila unit visual tidak punya bukti sel native, misalnya
    gambar tabel yang ditempel ke worksheet.
    """
    if survey is None:
        return True
    sheets = {sheet.name: sheet for sheet in survey.sheets}
    for region in visual_regions:
        if not region.cells:
            return True
        sheet = sheets.get(region.sheet_name)
        if sheet is None:
            return True
        if any(
            item.kind == "mixed"
            and (item.title or "").startswith("Image ")
            and _bounds_overlap(_region_bounds(item), _region_bounds(region))
            for item in sheet.regions
        ):
            return True
    return False


_AGENT_NATIVE_CONTEXT_FILE = "excel_native_agent.json"
_AGENT_VISUAL_INSTRUCTION = (
    "Workbook ini sudah diekstrak secara native: nilai sel, formula, dan data grafik "
    "tersimpan di SQLite (kolom row_role menandai subtotal/grand total). Fokuskan "
    "Markdown pada konteks visual: judul, anotasi, makna grafik, dan tata letak. "
    "Tabel angka tidak perlu ditranskripsi ulang; bila ditulis untuk keterbacaan, "
    "tabel tersebut tidak di-ingest ke SQLite dan angka native tetap menjadi acuan."
)


def _source_signature(source: Path) -> str:
    stat = source.stat()
    return f"{stat.st_size}:{stat.st_mtime_ns}"


def ensure_excel_native_artifacts(
    excel_path: str | Path,
    output_root: str | Path,
    *,
    settings: Settings | None = None,
) -> dict[str, Any] | None:
    """Siapkan artefak native Excel untuk Agent Mode (render per halaman).

    Survei dan penyimpanan SQLite hanya dijalankan sekali per versi file; panggilan
    berikutnya membaca cache ``regions/excel_native_agent.json``. Mengembalikan
    ``None`` bila survei native dinonaktifkan atau gagal, sehingga pemanggil tetap
    dapat melanjutkan alur visual lama.
    """
    source = Path(excel_path).resolve()
    root = Path(output_root).resolve()
    resolved_settings = settings or get_settings()
    if not resolved_settings.excel_native_survey:
        return None

    context_path = root / "regions" / _AGENT_NATIVE_CONTEXT_FILE
    db_path = root / "databases" / f"{source.stem}.sqlite"
    signature = _source_signature(source)
    if context_path.exists() and db_path.exists():
        try:
            cached = json.loads(context_path.read_text(encoding="utf-8"))
        except (OSError, ValueError):
            cached = None
        if isinstance(cached, dict) and cached.get("signature") == signature:
            return cached

    try:
        with tempfile.TemporaryDirectory(prefix="excel_agent_native_") as tmp:
            survey_source = source
            try:
                survey_source = _convert_workbook_to_xlsx_copy(
                    source, Path(tmp) / "survey"
                )
            except RuntimeError:
                if source.suffix.lower() in {".xls", ".ods"}:
                    raise
            survey = survey_excel_workbook(survey_source, settings=resolved_settings)
        survey.source_file = str(source)
        survey.workbook_format = source.suffix.lower().lstrip(".")
        visual_regions = plan_excel_visual_regions(survey, settings=resolved_settings)
        table_names = persist_excel_native_data(survey, db_path)
        artifacts = _describe_excel_native_artifacts(survey, db_path, table_names, [])
        native_markdown, _ = compose_excel_markdown(
            survey,
            fallback_title=source.stem,
            artifacts=artifacts,
        )
    except (OSError, ValueError, RuntimeError, TypeError, sqlite3.Error) as exc:
        logger.warning("[Excel] Artefak native Agent Mode gagal disiapkan: %s", exc)
        return None

    markdown_path = root / f"{source.stem}.native.md"
    markdown_path.write_text(native_markdown, encoding="utf-8")
    context: dict[str, Any] = {
        "signature": signature,
        "database_path": str(db_path),
        "native_markdown_path": str(markdown_path),
        "vlm_tables_needed": _visual_regions_need_vlm_tables(survey, visual_regions),
        "sheet_roles": {
            sheet.name: sheet.role for sheet in survey.sheets if sheet.visible
        },
        "native_tables": [artifact.model_dump(mode="json") for artifact in artifacts],
        "warnings": list(survey.warnings),
        "instruction": _AGENT_VISUAL_INSTRUCTION,
    }
    context_path.parent.mkdir(parents=True, exist_ok=True)
    context_path.write_text(
        json.dumps(context, ensure_ascii=False, indent=2), encoding="utf-8"
    )
    return context


def process_multipage_excel(
    excel_path: str | Path,
    pipeline: Any = None,
    *,
    llm: Any = None,
    output_dir: str | Path | None = None,
    dpi: int = DEFAULT_DPI,
    forced_specs: list[str] | str | None = None,
    forced_doc_type: str | None = None,
    db_path: str | Path | None = None,
    auto_tabular_db: bool = True,
    force_all_tables: bool = False,
    output_markdown_path: str | Path | None = None,
    resume: bool = False,
) -> ExtractedDocument:
    """
    Proses workbook Excel dengan jalur native dan visual yang adaptif.

    Dashboard atau blok campuran dirender untuk VLM, sedangkan ringkasan, formula,
    grafik native, dan data detail dipertahankan langsung dari struktur workbook.
    """
    path_obj = Path(excel_path).resolve()
    if not path_obj.exists():
        raise FileNotFoundError(f"File Excel tidak ditemukan: {path_obj}")

    with tempfile.TemporaryDirectory(prefix="excel_ingest_") as tmp:
        temporary_root = Path(tmp)
        settings = getattr(pipeline, "settings", None) or get_settings()
        survey: ExcelWorkbookSurvey | None = None
        visual_regions: list[ExcelRegion] = []
        survey_source = path_obj
        if settings.excel_native_survey:
            try:
                try:
                    survey_source = _convert_workbook_to_xlsx_copy(
                        path_obj, temporary_root / "survey"
                    )
                except RuntimeError:
                    if path_obj.suffix.lower() in {".xls", ".ods"}:
                        raise
                    survey_source = path_obj
                survey = survey_excel_workbook(survey_source, settings=settings)
                survey = split_excel_regions_into_tiles(survey, settings=settings)
                survey.source_file = str(path_obj)
                survey.workbook_format = path_obj.suffix.lower().lstrip(".")
                visual_regions = plan_excel_visual_regions(
                    survey, settings=settings
                )
                logger.info(
                    "[Excel] Survei native merencanakan %d unit visual pada %d sheet terlihat.",
                    len(visual_regions),
                    sum(sheet.visible for sheet in survey.sheets),
                )
            except (OSError, ValueError, RuntimeError, TypeError) as exc:
                logger.warning(
                    "[Excel] Survei struktur gagal; lanjutkan render seluruh sheet: %s",
                    exc,
                )

        pdf_path = temporary_root / f"{path_obj.stem}.pdf"
        region_rendered = False
        if settings.excel_region_rendering and visual_regions:
            region_rendered = _convert_excel_regions_to_pdf(
                path_obj,
                pdf_path,
                temporary_root / "lo_region_profile",
                visual_regions,
            )
        needs_visual_pipeline = bool(visual_regions) or survey is None
        if not region_rendered and needs_visual_pipeline:
            logger.info("[Excel] Mengonversi '%s' ke PDF sementara...", path_obj.name)
            pdf_path = convert_excel_to_pdf(path_obj, tmp)
        elif not needs_visual_pipeline:
            logger.info("[Excel] Tidak ada unit visual; ekstraksi diselesaikan secara native.")
        native_by_page = (
            {
                index: region.native_text
                for index, region in enumerate(visual_regions, start=1)
            }
            if region_rendered
            else None
        )
        dpi_by_page = (
            {
                index: region.render_dpi
                for index, region in enumerate(visual_regions, start=1)
            }
            if region_rendered
            else None
        )
        rescue_by_page = (
            {
                index: region.requires_vlm_reading
                for index, region in enumerate(visual_regions, start=1)
            }
            if region_rendered
            else None
        )
        if needs_visual_pipeline:
            logger.info(
                "[Excel] Konversi selesai, lanjut render unit visual dan ekstraksi VLM."
            )
            result = process_multipage_pdf(
                pdf_path=pdf_path,
                pipeline=pipeline,
                llm=llm,
                output_dir=output_dir,
                dpi=dpi,
                forced_specs=forced_specs,
                forced_doc_type=forced_doc_type,
                db_path=db_path,
                auto_tabular_db=auto_tabular_db
                and _visual_regions_need_vlm_tables(survey, visual_regions),
                force_all_tables=force_all_tables,
                output_markdown_path=None,
                source_file_for_records=path_obj,
                resume=resume,
                native_text_by_page_override=native_by_page,
                dpi_by_page_override=dpi_by_page,
                force_vlm_reading_by_page_override=rescue_by_page,
            )
        else:
            result = ExtractedDocument(
                source_file=str(path_obj),
                title=path_obj.stem,
                doc_type="spreadsheet",
                full_markdown="",
                total_pages=1,
            )

        if survey is not None:
            result.excel_workbook = survey
            if output_dir:
                output_root = Path(output_dir).resolve().parent
            elif output_markdown_path:
                output_root = Path(output_markdown_path).resolve().parent
            else:
                output_root = Path("output") / path_obj.stem
            manifest_path = output_root / "regions" / "excel_manifest.json"
            manifest_path.parent.mkdir(parents=True, exist_ok=True)
            manifest_path.write_text(
                json.dumps(survey.model_dump(mode="json"), ensure_ascii=False, indent=2),
                encoding="utf-8",
            )
            try:
                render_excel_sheet_previews(path_obj, survey, output_root, temporary_root)
            except Exception as exc:  # noqa: BLE001
                logger.warning("[Excel] Gambar pembanding per sheet gagal: %s", exc)

            if db_path:
                native_db_path = Path(db_path).resolve()
            elif output_markdown_path:
                native_db_path = (
                    Path(output_markdown_path).resolve().parent
                    / "databases"
                    / f"{path_obj.stem}.sqlite"
                )
            elif output_dir:
                native_db_path = (
                    Path(output_dir).resolve()
                    / "databases"
                    / f"{path_obj.stem}.sqlite"
                )
            else:
                native_db_path = (
                    Path("output")
                    / path_obj.stem
                    / "databases"
                    / f"{path_obj.stem}.sqlite"
                )
            if auto_tabular_db:
                result.excel_native_tables = persist_excel_native_data(
                    survey, native_db_path
                )
                csv_paths: list[Path] = []
                try:
                    from .tabular_db import TabularDatabaseManager

                    csv_dir = (
                        native_db_path.parent.parent / "csv"
                        if native_db_path.parent.name == "databases"
                        else native_db_path.parent / "csv"
                    )
                    csv_paths = TabularDatabaseManager(native_db_path).export_to_csv(
                        output_dir=csv_dir
                    )
                except Exception as exc:  # noqa: BLE001
                    logger.warning("[Excel] Ekspor data native ke CSV gagal: %s", exc)
                result.excel_native_artifacts = _describe_excel_native_artifacts(
                    survey,
                    native_db_path,
                    result.excel_native_tables,
                    csv_paths,
                )

            visual_by_sheet: dict[str, str] = {}
            if region_rendered and len(result.pages) == len(visual_regions):
                tile_groups: dict[str, list[ExcelRegion]] = {}
                for region in visual_regions:
                    if region.parent_region_id:
                        tile_groups.setdefault(region.parent_region_id, []).append(region)
                rendered_tile_groups: set[str] = set()
                for page, region in zip(result.pages, visual_regions, strict=True):
                    if region.parent_region_id:
                        parent_id = region.parent_region_id
                        if parent_id in rendered_tile_groups:
                            continue
                        rendered_tile_groups.add(parent_id)
                        combined = _combine_tile_regions(tile_groups[parent_id])
                        content = _native_region_markdown(combined)
                        sheet_name = combined.sheet_name
                    else:
                        content = page.markdown_content.strip()
                        sheet_name = region.sheet_name
                    if content:
                        visual_by_sheet[sheet_name] = "\n\n".join(
                            part
                            for part in (
                                visual_by_sheet.get(sheet_name),
                                content,
                            )
                            if part
                        )
            result.full_markdown, _ = compose_excel_markdown(
                survey,
                fallback_title=result.title or path_obj.stem,
                artifacts=result.excel_native_artifacts,
                visual_by_sheet=visual_by_sheet,
            )
            result.title = _workbook_title(survey, result.title or path_obj.stem)
            result.total_tables = sum(
                region.kind == "table"
                for sheet in survey.sheets
                for region in sheet.regions
                if sheet.visible and sheet.role != "support"
            )
            if output_markdown_path:
                markdown_target = Path(output_markdown_path).resolve()
                markdown_target.parent.mkdir(parents=True, exist_ok=True)
                markdown_target.write_text(result.full_markdown, encoding="utf-8")

        return result


__all__ = [
    "compose_excel_markdown",
    "convert_excel_for_extraction",
    "convert_excel_to_pdf",
    "excel_page_count",
    "persist_excel_native_data",
    "plan_excel_visual_regions",
    "process_multipage_excel",
    "split_excel_regions_into_tiles",
    "survey_excel_workbook",
]
